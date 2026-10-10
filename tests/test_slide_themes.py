"""主题驱动的 PPT：各版式、封面与章节页变体、东亚字体、页脚、表格、提醒；以及 LibreOffice 的字体替换表。"""

import json
from xml.dom import minidom

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from acks_office import fonts, themes, utils
from acks_office.cli import main
from acks_office.pptx import create_pptx
from acks_office.render.slides import NO_STYLE_TABLE

DECK = [
    {"layout": "title", "kicker": "2026 年第三季度", "title": "经营回顾", "subtitle": "一个季度的关键变化",
     "footer": "经营分析部 · 2026 年 10 月",
     "kpis": [{"label": "营收", "value": "1.28", "unit": "亿元"}, {"label": "同比", "value": "+24%", "highlight": True}]},
    {"layout": "section", "number": "01", "kicker": "PART ONE · 第一部分", "title": "生意", "title_en": "The Business",
     "subtitle": "一个季度的账本。"},
    {"layout": "data", "title": "营收增长 24%", "subtitle": "华东贡献近一半",
     "kpis": [{"label": "营收", "value": "1.28", "unit": "亿元", "delta": "同比 +24.0%"}],
     "chart": {"title": "分区域营收占比", "unit": "单位：%", "data": [["华东", 46], ["华南", 27]], "highlight": 0,
               "note": "示例数据"}},
    {"layout": "table", "title": "三件事",
     "table": {"header": ["序号", "事项", "负责"], "rows": [["01", "华东加密", "拓展部"], ["02", "稳住成本", "供应链部"]]}},
    {"layout": "number", "label": "营收 · REVENUE", "value": "1.28", "unit": "亿元", "delta_label": "同比",
     "delta": "+24%", "text": "华东贡献近一半。", "chart": {"data": [["华东", 46], ["其他", 54]], "highlight": 0}},
    {"layout": "quote", "quote": "好咖啡不必昂贵，但一定要认真。", "author": "林舟", "role": "创始人"},
    {"layout": "content", "title": "要点", "content": "第一点\n- 第二点", "notes": "讲者备注"},
]


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    monkeypatch.setenv("ACKS_OFFICE_HOME", str(tmp_path / "home"))
    themes.clear_cache()
    yield
    themes.clear_cache()


def _deck(tmp_path, theme, slides=DECK, **kwargs):
    out = tmp_path / f"{theme}.pptx"
    kwargs.setdefault("brand_name", "栖木咖啡")
    result = create_pptx("2026 Q3 经营回顾", slides, str(out), theme=theme, **kwargs)
    return out, result, Presentation(str(out))


def _texts(slide):
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]
    for shape in slide.shapes:
        if shape.has_table:
            texts += [cell.text for row in shape.table.rows for cell in row.cells]
    return texts


def _background(slide):
    return str(slide.background.fill.fore_color.rgb)


@pytest.mark.parametrize("theme", ["neutral", "slate", "folio"])
def test_every_layout_renders(tmp_path, theme):
    out, result, prs = _deck(tmp_path, theme)
    assert result["slides_count"] == len(DECK) and result["theme"] == theme
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000  # 16:9
    flat = "\n".join(t for slide in prs.slides for t in _texts(slide))
    for text in ("经营回顾", "生意", "营收增长 24%", "华东加密", "1.28", "好咖啡不必昂贵", "第二点"):
        assert text in flat
    assert prs.slides[6].notes_slide.notes_text_frame.text == "讲者备注"


def test_runs_carry_east_asian_fonts(tmp_path, monkeypatch):
    monkeypatch.setattr(fonts, "family_installed", lambda name: False)
    out, _, prs = _deck(tmp_path, "slate")
    run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    rpr = run._r.find(qn("a:rPr"))
    assert rpr.find(qn("a:latin")).get("typeface") == "Source Sans 3"
    assert rpr.find(qn("a:ea")).get("typeface") == "Noto Sans SC"


def test_cover_and_section_variants(tmp_path):
    _, _, slate = _deck(tmp_path, "slate")
    cover = slate.slides[0]
    assert "+24%" in _texts(cover) and "栖木咖啡" in _texts(cover)  # 右侧关键数字栏与品牌行
    assert any(shape.shape_type == 1 and str(shape.fill.fore_color.rgb) == "F5F7FA" for shape in cover.shapes)
    assert _background(slate.slides[1]) == "FFFFFF"

    _, _, folio = _deck(tmp_path, "folio", publication="栖木咖啡季度刊", issue="07", season="Autumn 2026",
                        date="2026.10")
    assert _background(folio.slides[0]) == "F3F1EC"
    cover_text = "\n".join(_texts(folio.slides[0]))
    assert "ISSUE 07 — AUTUMN 2026" in cover_text and "07" in cover_text and "2026.10" in cover_text
    section = folio.slides[1]
    assert _background(section) == "1F2A44" and "生意  The Business" in _texts(section)


def test_footer_uses_theme_template_and_page_numbers(tmp_path):
    _, _, slate = _deck(tmp_path, "slate", short_title="Q3 经营回顾")
    texts = _texts(slate.slides[2])
    assert "栖木咖啡 · Q3 经营回顾" in texts and "03" in texts

    _, _, folio = _deck(tmp_path, "folio", publication="季度刊", issue="07")
    assert "季度刊 · 第 07 期" in _texts(folio.slides[2])

    _, _, custom = _deck(tmp_path, "slate", footer_label="内部资料")
    assert "内部资料" in _texts(custom.slides[2])


def test_table_slide_uses_native_table_without_style(tmp_path):
    _, _, prs = _deck(tmp_path, "slate")
    (graphic,) = [s for s in prs.slides[3].shapes if s.has_table]
    table = graphic.table
    assert table._tbl.tblPr.find(qn("a:tableStyleId")).text == NO_STYLE_TABLE
    assert [c.text for c in table.rows[0].cells] == ["序号", "事项", "负责"]
    number_run = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
    assert str(number_run.font.color.rgb) == "1F5FAE"  # 序号用强调色
    assert "a:lnB" in table.cell(1, 0)._tc.xml


def test_long_titles_shrink_to_fit(tmp_path):
    long = "这是一个非常非常长的标题，用来检查标题会不会自动缩小字号以免超出版面" * 2
    _, _, prs = _deck(tmp_path, "slate", slides=[{"layout": "content", "title": long, "content": "x"}])
    size = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt
    assert size < themes.load_theme("slate").size("type.slide.title")


def test_images_unknown_layouts_and_legacy_content(tmp_path, png):
    png(400, 200, name="pic.png")
    slides = [{"layout": "image", "image": "pic.png", "caption": "门店外景"},
              {"layout": "image", "image": "missing.png"},
              {"layout": "fancy", "title": "未知版式", "content": "• 要点"},
              {"title": "旧写法", "content": "第一行\n第二行"}]
    _, result, prs = _deck(tmp_path, "neutral", slides=slides, base_dir=str(tmp_path))
    assert any(s.shape_type == 13 for s in prs.slides[0].shapes)  # 图片
    codes = [w["code"] for w in result["warnings"]]
    assert "IMAGE_NOT_FOUND" in codes and "UNKNOWN_LAYOUT" in codes
    assert "• 第一行\n• 第二行" in _texts(prs.slides[3])


def test_cli_slides_file_with_deck_meta(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    deck = {"meta": {"title": "经营回顾特辑", "publication": "季度刊", "issue": "07"}, "slides": DECK[:2]}
    (tmp_path / "deck.json").write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
    code = main(["create", "pptx", "-o", "刊物.pptx", "--slides-file", "deck.json", "--theme", "folio", "--json"])
    env = json.loads(capsys.readouterr().out)
    assert code == 0 and env["data"]["theme"] == "folio" and env["data"]["slides_count"] == 2
    prs = Presentation(str(tmp_path / "刊物.pptx"))
    assert prs.core_properties.title == "经营回顾特辑"
    assert "季度刊　　ISSUE 07" in "\n".join(_texts(prs.slides[0]))
    (tmp_path / "bad.json").write_text(json.dumps({"slides": []}), encoding="utf-8")
    assert main(["create", "pptx", "-o", "x.pptx", "--slides-file", "bad.json", "--json"]) == 1
    assert json.loads(capsys.readouterr().out)["error"]["code"] == "INVALID_INPUT"


def test_default_slide_theme_is_neutral(tmp_path):
    plain = tmp_path / "plain.pptx"
    assert create_pptx("T", [{"title": "经营回顾", "layout": "title"}], str(plain))["theme"] == "neutral"
    assert "ACKS" not in "\n".join(_texts(Presentation(str(plain)).slides[0]))


# ---------------------------------------------------------------- LibreOffice 字体替换表

def _face(families, localized):
    return fonts.FontFace("/x.ttc", 0, tuple(families), "Regular", 400, False, False, False, tuple(localized))


def test_localized_family_names_follow_ui_language(monkeypatch):
    faces = [_face(["PingFang SC", "苹方-简", "蘋方-簡"],
                   [("en", "PingFang SC"), ("zh-Hans", "苹方-简"), ("zh-Hant", "蘋方-簡")]),
             _face(["Helvetica Neue"], [("en", "Helvetica Neue")])]
    monkeypatch.setattr(fonts, "font_faces", lambda write_cache=True: faces)
    assert fonts.localized_family_names("zh-Hans-CN") == {"PingFang SC": "苹方-简"}
    assert fonts.localized_family_names("zh-Hant-TW") == {"PingFang SC": "蘋方-簡"}
    assert fonts.localized_family_names("en-US") == {}


def test_seed_profile_writes_valid_substitution_table(monkeypatch, tmp_path):
    monkeypatch.setattr(fonts, "localized_family_names",
                        lambda ui_language=None: {"PingFang SC": "苹方-简", "A & B": "甲<乙>", "Bad\x00": "坏"})
    count = utils._seed_profile(tmp_path / "profile")
    xml = (tmp_path / "profile" / "user" / "registrymodifications.xcu").read_bytes()
    minidom.parseString(xml)
    assert count == 2 and "苹方-简".encode() in xml and b"A &amp; B" in xml and b"Bad" not in xml

    monkeypatch.setattr(fonts, "localized_family_names", lambda ui_language=None: {})
    assert utils._seed_profile(tmp_path / "empty") == 0 and not (tmp_path / "empty").exists()


def test_name_table_records_localized_family_names(tmp_path):
    pytest.importorskip("fontTools")
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    fb = FontBuilder(1000, isTTF=True)
    fb.setupGlyphOrder([".notdef", "uni4E2D"])
    fb.setupCharacterMap({0x4E2D: "uni4E2D"})
    pen = TTGlyphPen(None)
    pen.moveTo((0, 0)); pen.lineTo((0, 500)); pen.lineTo((500, 500)); pen.closePath()
    fb.setupGlyf({".notdef": pen.glyph(), "uni4E2D": pen.glyph()})
    fb.setupHorizontalMetrics({".notdef": (1000, 0), "uni4E2D": (1000, 0)})
    fb.setupHorizontalHeader(ascent=880, descent=-120)
    # 实际字体用 Windows 语言编号记录本地化名称；fontTools 里简体中文（0x0804）写作 zh
    fb.setupNameTable({"familyName": {"en": "Acks Hei", "zh": "测试黑体"}, "styleName": "Regular",
                       "psName": "AcksHei-Regular"})
    fb.setupOS2()
    fb.setupPost()
    path = tmp_path / "AcksHei.ttf"
    fb.save(str(path))
    (face,) = fonts.read_font_faces(str(path))
    assert face.families[0] == "Acks Hei" and "测试黑体" in face.families
    assert ("en", "Acks Hei") in face.localized and ("zh-Hans", "测试黑体") in face.localized
