"""Word：Markdown 呈现、品牌与页脚、表格提取、合并。"""

import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from acks_office.docx import create_word, extract_text, merge_documents

TABLE = "| 区域 | 营收 |\n|---|--:|\n| 华东 | 5,888 |\n| 华南 | 3,456 |"


def _xml(path, prefix="word/"):
    """docx 里所有 XML 部件的文本（正文、页眉、页脚等）。"""
    with zipfile.ZipFile(path) as z:
        return "".join(z.read(n).decode("utf-8") for n in z.namelist()
                       if n.startswith(prefix) and n.endswith(".xml"))


def _runs(doc):
    for p in doc.paragraphs:
        for r in p.iter_inner_content():
            yield from getattr(r, "runs", [r])  # 链接里的文字在 Hyperlink.runs 里


@pytest.mark.parametrize("theme", ["default", "folio"])
def test_inline_markdown_becomes_real_formatting(tmp_path, theme):
    out = tmp_path / "a.docx"
    create_word("标题", "普通 **粗体** *斜体* ~~删除~~ [链接](https://example.com/?a=1&b=2)", str(out), theme=theme)

    doc = Document(str(out))
    runs = {r.text: r for r in _runs(doc) if r.text.strip()}
    assert runs["粗体"].bold and runs["斜体"].italic and runs["删除"].font.strike
    assert not any("**" in r.text for r in _runs(doc))

    (link,) = [h for p in doc.paragraphs for h in p.hyperlinks]
    assert link.text == "链接" and link.address == "https://example.com/?a=1&b=2"


@pytest.mark.parametrize("theme", ["default", "folio"])
def test_table_is_a_word_table_and_extracts_by_row(tmp_path, theme):
    out = tmp_path / "t.docx"
    create_word("标题", "前言\n\n" + TABLE + "\n\n结语", str(out), theme=theme)

    doc = Document(str(out))
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["区域", "营收"]
    lines = extract_text(str(out)).splitlines()
    assert lines.index("区域 | 营收") < lines.index("华东 | 5,888") < lines.index("结语")


def test_line_breaks_are_kept(tmp_path):
    out = tmp_path / "b.docx"
    create_word("标题", "第一行\n第二行", str(out), theme="default")
    assert "第一行\n第二行" in extract_text(str(out))


def test_cover_shows_title_once_with_subtitle(tmp_path):
    out = tmp_path / "c.docx"
    create_word("季度报告", "正文", str(out), subtitle="营收与会员")
    texts = [p.text for p in Document(str(out)).paragraphs]
    assert sum("季度报告" in t for t in texts) == 1
    assert "营收与会员" in texts


def test_empty_brand_removes_all_brand_text(tmp_path):
    out = tmp_path / "d.docx"
    create_word("季度报告", "正文", str(out), brand_name="")
    xml = _xml(out)
    assert "ACKS" not in xml.upper()
    assert "季度报告" in _xml(out, "word/footer")  # 无品牌时页脚显示标题


def test_custom_brand_replaces_default(tmp_path):
    out = tmp_path / "e.docx"
    create_word("季度报告", "正文", str(out), brand_name="栖木咖啡")
    xml = _xml(out)
    assert "栖木咖啡" in xml and "ACKS" not in xml.upper()


def _num_ids(doc):
    ids = []
    for p in doc.paragraphs:
        num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
        if num_pr is not None and num_pr.numId is not None:
            ids.append((p.text, num_pr.numId.val))
    return ids


def test_default_theme_restarts_each_ordered_list(tmp_path):
    out = tmp_path / "n.docx"
    create_word("标题", "1. 一\n2. 二\n\n段落\n\n1. 三\n2. 四\n\n段落\n\n3. 从三开始", str(out), theme="default")

    doc = Document(str(out))
    ids = dict(_num_ids(doc))
    assert ids["一"] == ids["二"] != ids["三"] == ids["四"] != ids["从三开始"]
    numbering = doc.part.numbering_part.element
    starts = {text: numbering.num_having_numId(num_id).find(f".//{qn('w:startOverride')}").get(qn("w:val"))
              for text, num_id in ids.items()}
    assert starts == {"一": "1", "二": "1", "三": "1", "四": "1", "从三开始": "3"}


def test_task_items_have_only_a_checkbox(tmp_path):
    out = tmp_path / "k.docx"
    create_word("标题", "- [ ] 待办\n- [x] 完成", str(out), theme="default")  # default 即 neutral
    items = [p for p in Document(str(out)).paragraphs if p.text.endswith(("待办", "完成"))]
    assert [p.text for p in items] == ["☐ 待办", "☑ 完成"]
    assert all(p._p.pPr.numPr is None for p in items)


@pytest.mark.parametrize("size, expect_native", [((120, 60), True), ((4000, 1000), False)])
def test_images_are_never_upscaled(tmp_path, png, size, expect_native):
    png(*size, name="pic.png")
    out = tmp_path / "i.docx"
    create_word("标题", "![图注](pic.png)", str(out), theme="default", base_dir=str(tmp_path))

    doc = Document(str(out))
    (shape,) = doc.inline_shapes
    section = doc.sections[-1]
    text_width = section.page_width - section.left_margin - section.right_margin
    assert shape.width <= text_width
    assert (shape.width < text_width) == expect_native
    assert abs(shape.width / shape.height - size[0] / size[1]) < 0.01
    assert "图注" in extract_text(str(out))


def test_missing_image_becomes_placeholder_with_warning(tmp_path):
    result = create_word("标题", "![图注](nope.png)", str(tmp_path / "m.docx"), base_dir=str(tmp_path))
    assert "IMAGE_NOT_FOUND" in [w["code"] for w in result["warnings"]]
    assert "[图片：图注]" in extract_text(str(tmp_path / "m.docx"))


def test_merge_keeps_images_and_refuses_missing_files(tmp_path, png):
    image = png(200, 100, name="pic.png")
    first, second, merged = tmp_path / "1.docx", tmp_path / "2.docx", tmp_path / "merged.docx"
    create_word("第一份", "第一份正文", str(first), theme="default")
    create_word("第二份", "第二份正文\n\n![图](pic.png)\n\n[链接](https://example.com)", str(second),
                theme="default", base_dir=str(tmp_path))

    with pytest.raises(FileNotFoundError):  # 3.0：缺文件直接报错，不会只合并一部分
        merge_documents([str(first), str(tmp_path / "missing.docx"), str(second)], str(merged))
    assert not merged.exists()
    result = merge_documents([str(first), str(second)], str(merged))

    assert result["merged_count"] == 2 and "skipped" not in result
    doc = Document(str(merged))
    text = extract_text(str(merged))
    assert text.index("第一份正文") < text.index("第二份正文")
    blips = doc.element.body.findall(".//" + qn("a:blip"))
    assert len(blips) == 1
    assert doc.part.rels[blips[0].get(qn("r:embed"))].target_part.blob == image.read_bytes()
    assert [h.address for p in doc.paragraphs for h in p.hyperlinks] == ["https://example.com"]
    assert doc.element.body.find(".//" + qn("w:br")).get(qn("w:type")) == "page"


def test_merge_warns_about_list_numbering(tmp_path):
    first, second = tmp_path / "1.docx", tmp_path / "2.docx"
    create_word("一", "正文", str(first), theme="default")
    create_word("二", "1. 甲\n2. 乙", str(second), theme="default")
    result = merge_documents([str(first), str(second)], str(tmp_path / "m.docx"))
    assert {w["code"] for w in result["warnings"]} == {"MERGE_NUMBERING"}
