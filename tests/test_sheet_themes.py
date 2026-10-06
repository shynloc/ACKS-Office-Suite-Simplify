"""主题驱动的 Excel：标题行、表头变体、数字格式与右边距、合计行、重点单元格、多张工作表与图表、读取时识别标题行。"""

import json
import zipfile

import openpyxl
import pytest

from acks_office import themes
from acks_office.cli import main
from acks_office.render.common import records_to_rows
from acks_office.render.sheet import number_format, padded, sheet_name
from acks_office.themes import validate_theme
from acks_office.xlsx import create_excel, extract_data

DATA = [["区域", "7 月", "8 月", "同比"],
        ["华东", 1846, 1962, 0.258],
        ["华南", 1092, -1148.5, 0.22],
        ["合计", 2938, 813.5, 0.24]]


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    monkeypatch.setenv("ACKS_OFFICE_HOME", str(tmp_path / "home"))
    themes.clear_cache()
    yield
    themes.clear_cache()


def _sheet(path, index=0):
    return openpyxl.load_workbook(path).worksheets[index]


def _rgb(color):
    return color.rgb[-6:].upper() if color is not None and isinstance(color.rgb, str) else None


def _sides(cell):
    sides = {side: getattr(cell.border, side) for side in ("left", "right", "top", "bottom")}
    return {side: value.style if value is not None else None for side, value in sides.items()}


def test_number_formats_and_padding():
    assert number_format([1, 2, None, "x"]) == "#,##0_);(#,##0);#,##0_);@_)"
    assert number_format([1.5, 2]) == "#,##0.0_);(#,##0.0);#,##0.0_);@_)"
    assert number_format([1.23456, float("nan"), True]) == "#,##0.00_);(#,##0.00);#,##0.00_);@_)"
    assert number_format(["a", True, None]) is None
    assert padded("0.0%") == "0.0%_0"
    assert padded("0;;") == "0_0;;"  # 空段保持为空，仍然不显示
    assert padded('0" ;件";-0') == '0" ;件"_0;-0_0'  # 引号里的分号不拆
    assert records_to_rows([{"a": 1}, {"b": 2, "a": 3}]) == [["a", "b"], [1, None], [3, 2]]


def test_sheet_names_are_valid_and_unique():
    used = []
    assert sheet_name("Q3: 营收/成本?", used) == "Q3 营收成本"
    assert sheet_name("q3 营收成本", used) == "q3 营收成本 (2)"  # Excel 的工作表名不分大小写
    assert len(sheet_name("长" * 40, used)) == 31
    assert sheet_name("", used) == "Sheet"


def test_neutral_title_row_header_fill_formats_and_total(tmp_path):
    out = tmp_path / "表.xlsx"
    result = create_excel("Q3 分区域营收", DATA, str(out), theme="neutral")
    theme = themes.load_theme("neutral")
    assert result["theme"] == "neutral" and (result["rows"], result["columns"]) == (3, 4)
    ws = _sheet(out)
    assert ws.title == "Q3 分区域营收" and ws["A1"].value == "Q3 分区域营收"
    head = ws["A2"]
    assert head.value == "区域" and head.font.b and _rgb(head.fill.fgColor) == theme.hex("surface")
    assert head.border.bottom.style == "medium" and _rgb(head.border.bottom.color) == theme.hex("rule_mid")
    assert ws.freeze_panes == "A3" and ws.sheet_view.showGridLines is False
    assert _rgb(ws.sheet_properties.tabColor) == theme.hex("accent")
    assert ws.print_title_rows == "$2:$2"
    # 右对齐的右边距写在数字格式里，表头和数字的右边缘一致；左对齐的文字用缩进
    assert ws["B3"].number_format == "#,##0_)_0;(#,##0)_0;#,##0_)_0;@_)_0"
    assert ws["C4"].number_format.startswith("#,##0.0_)")
    assert ws["B2"].number_format == "@_)_0"
    assert (ws["B3"].alignment.horizontal, ws["B3"].alignment.indent) == ("right", 0)
    assert (ws["A3"].alignment.horizontal, ws["A3"].alignment.indent) == ("left", 1)
    # 合计行：加粗，上下各一条细线；它上面一行不再画线
    assert ws["A5"].value == "合计" and ws["B5"].font.b
    assert _sides(ws["B5"])["top"] == "thin" and _sides(ws["B5"])["bottom"] == "thin"
    assert _rgb(ws["B5"].border.top.color) == theme.hex("rule_strong")
    assert _sides(ws["B4"])["bottom"] is None and _sides(ws["B3"])["bottom"] == "thin"
    # 负数用主题的负数色（条件格式）
    rules = [rule for cf in ws.conditional_formatting for rule in cf.rules]
    assert rules and {rule.operator for rule in rules} == {"lessThan"}
    assert _rgb(rules[0].dxf.font.color) == theme.hex("negative")


def test_without_title_the_header_is_row_one(tmp_path):
    out = tmp_path / "a.xlsx"
    result = create_excel(None, DATA, str(out), theme="neutral", sheet_name="明细")
    ws = _sheet(out)
    assert result["sheets"] == [{"name": "明细", "rows": 3, "columns": 4}]
    assert ws.title == "明细" and ws["A1"].value == "区域" and ws.freeze_panes == "A2"
    assert extract_data(str(out))[0] == {"区域": "华东", "7 月": 1846, "8 月": 1962, "同比": 0.258}


def test_folio_label_header_total_bar_and_ink_highlight(tmp_path):
    out = tmp_path / "folio.xlsx"
    spec = {"name": "Q3", "title": "Q3 营收", "header": DATA[0], "rows": DATA[1:], "highlight": [[0, 1]]}
    create_excel("Q3", None, str(out), theme="folio", sheets=[spec])
    theme = themes.load_theme("folio")
    ws = _sheet(out)
    head = ws["B2"]
    assert head.fill.fill_type is None and not head.font.b and _rgb(head.font.color) == theme.hex("muted")
    assert head.border.bottom.style == "medium" and _rgb(head.border.bottom.color) == theme.hex("rule_mid")
    total = ws["B5"]  # bar：上方一条强调色粗线，下方不画
    assert total.border.top.style == "medium" and _rgb(total.border.top.color) == theme.hex("accent")
    assert _sides(total)["bottom"] is None
    assert set(_sides(ws["B3"]).values()) == {"medium"} and _rgb(ws["B3"].border.left.color) == theme.hex("ink")
    assert ws["A3"].font.name == theme.font("body").cn and ws["A1"].font.name == theme.font("title").cn


def test_slate_total_rules_and_accent_highlight(tmp_path):
    out = tmp_path / "slate.xlsx"
    create_excel("Q3", None, str(out), theme="slate",
                 sheets=[{"title": "Q3", "header": DATA[0], "rows": DATA[1:], "highlight": [[1, 2]]}])
    theme = themes.load_theme("slate")
    ws = _sheet(out)
    assert _sides(ws["C5"])["top"] == "thin" and _sides(ws["C5"])["bottom"] == "thin"
    assert set(_sides(ws["C4"]).values()) == {"medium"} and _rgb(ws["C4"].border.top.color) == theme.hex("accent")


def test_appended_total_has_formulas_and_readable_values(tmp_path):
    out = tmp_path / "t.xlsx"
    spec = {"header": ["年份", "营收", "备注"], "rows": [[2024, 1200.5, "a"], [2025, 1400.25, None]], "total": True}
    create_excel(None, None, str(out), theme="neutral", sheets=[spec])
    ws = _sheet(out)
    assert ws["A4"].value == "合计" and ws["B4"].value == "=SUM(B2:B3)" and ws["C4"].value is None
    assert ws["A2"].number_format == "0" and ws["A2"].alignment.horizontal == "left"  # 年份不加千位分隔
    assert openpyxl.load_workbook(out, data_only=True).active["B4"].value == 2600.75  # 合计结果已写进文件
    assert extract_data(str(out))[-1] == {"年份": "合计", "营收": 2600.75, "备注": None}

    # total: false 时最后一行即使叫「合计」也按普通行处理
    plain = tmp_path / "p.xlsx"
    create_excel(None, None, str(plain), theme="neutral",
                 sheets=[{"header": ["区域", "数量"], "rows": [["一部", 1], ["合计", 1]], "total": False}])
    ws = _sheet(plain)
    assert not ws["B3"].font.b and _sides(ws["B3"])["top"] is None


def test_note_chart_and_multiple_sheets(tmp_path):
    out = tmp_path / "m.xlsx"
    sheets = [{"name": "Q3", "title": "Q3", "header": ["月份", "营收", "毛利"],
               "rows": [["7 月", 4036, 2381.8], ["8 月", 4262, 2472]], "note": "注：单位万元", "chart": {"values": [1, 2]}},
              {"name": "q3", "header": ["说明"], "rows": [["无数字"]], "chart": {}}]
    result = create_excel("经营数据", None, str(out), theme="slate", sheets=sheets, author="经营分析部")
    assert [s["name"] for s in result["sheets"]] == ["Q3", "q3 (2)"]
    assert [w["code"] for w in result["warnings"]].count("CHART_SKIPPED") == 1
    wb = openpyxl.load_workbook(out)
    assert wb.properties.title == "经营数据" and wb.properties.creator == "经营分析部"
    ws = wb.worksheets[0]
    assert ws["A5"].value == "注：单位万元" and ws["A5"].alignment.vertical == "bottom"  # 紧接表格，用行高留间距
    with zipfile.ZipFile(out) as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        xml = z.read(charts[0]).decode("utf-8")
    assert len(charts) == 1
    colors = [c.lstrip("#").upper() for c in themes.load_theme("slate").chart_colors()[:2]]
    assert all(f'val="{c}"' in xml for c in colors) and "<legend>" in xml
    assert extract_data(str(out))[-1]["月份"] == "注：单位万元"  # 说明紧接表格，不会多出空行


def test_text_numbers_align_right_but_are_not_summed_or_charted(tmp_path):
    out = tmp_path / "s.xlsx"
    spec = {"header": ["区域", "同比", "营收"], "rows": [["华东", "+25.8%", 100], ["华南", "—", 50]],
            "total": True, "chart": {}}
    result = create_excel(None, None, str(out), theme="neutral", sheets=[spec])
    ws = _sheet(out)
    assert ws["B2"].alignment.horizontal == "right" and ws["B2"].number_format.endswith(";@_0")
    assert ws["B4"].value is None and ws["C4"].value == "=SUM(C2:C3)"
    assert "CHART_SKIPPED" not in [w["code"] for w in result.get("warnings", [])]
    with zipfile.ZipFile(out) as z:
        xml = z.read("xl/charts/chart1.xml").decode("utf-8")
    assert "'Sheet1'!$C$2:$C$3" in xml  # 数值取营收列，不含合计行
    assert 'formatCode="#,##0_);(#,##0);#,##0_);@_)"' in xml  # 坐标轴不带右边距


def test_values_are_cleaned_and_bad_highlights_warn(tmp_path):
    out = tmp_path / "v.xlsx"
    spec = {"header": ["名称", "数值", "标签"],
            "rows": [["a\x07b", float("nan"), ["x", "y"]], ["c", 2, {"k": 1}]], "highlight": [[9, 9], "bad"]}
    result = create_excel(None, None, str(out), theme="neutral", sheets=[spec])
    ws = _sheet(out)
    assert ws["A2"].value == "ab" and ws["B2"].value is None
    assert ws["C2"].value == '["x", "y"]' and ws["C3"].value == '{"k": 1}'
    assert [w["code"] for w in result["warnings"]].count("BAD_HIGHLIGHT") == 2
    with pytest.raises(ValueError):
        create_excel(None, None, str(tmp_path / "e.xlsx"), theme="neutral", sheets=[])
    with pytest.raises(ValueError):
        create_excel(None, None, str(tmp_path / "f.xlsx"), theme="acks", sheets=[spec])


def test_extract_detects_a_title_row(tmp_path):
    titled = tmp_path / "titled.xlsx"
    wb = openpyxl.Workbook()
    for row in (["报表标题"], ["名称", "数量"], ["甲", 1]):
        wb.active.append(row)
    wb.save(titled)
    assert extract_data(str(titled)) == [{"名称": "甲", "数量": 1}]
    assert extract_data(str(titled), header=0)[0] == {"报表标题": "名称", "Unnamed: 1": "数量"}

    single = tmp_path / "single.xlsx"  # 只有一列的表格：第一行就是表头
    wb = openpyxl.Workbook()
    for row in (["名称"], ["甲"], ["乙"]):
        wb.active.append(row)
    wb.save(single)
    assert extract_data(str(single)) == [{"名称": "甲"}, {"名称": "乙"}]


def test_cli_records_sheets_and_legacy_limits(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    records = [{"名称": "甲", "数量": 1}, {"名称": "乙", "单价": 2.5}]
    (tmp_path / "r.json").write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
    assert main(["create", "xlsx", "-o", "r.xlsx", "--data-file", "r.json", "--theme", "neutral", "--json"]) == 0
    env = json.loads(capsys.readouterr().out)
    assert env["data"]["sheets"] == [{"name": "r", "rows": 2, "columns": 3}]
    assert [c.value for c in _sheet(tmp_path / "r.xlsx")[1]] == ["名称", "数量", "单价"]

    book = {"meta": {"title": "经营数据", "author": "经营分析部"},
            "sheets": [{"name": "A", "header": ["x", "y"], "rows": [["a", 1]]}]}
    (tmp_path / "s.json").write_text(json.dumps(book, ensure_ascii=False), encoding="utf-8")
    assert main(["create", "xlsx", "-o", "s.xlsx", "--data-file", "s.json", "--theme", "slate", "--json"]) == 0
    env = json.loads(capsys.readouterr().out)
    assert env["data"]["theme"] == "slate" and env["data"]["sheets"][0]["name"] == "A"
    assert openpyxl.load_workbook(tmp_path / "s.xlsx").properties.title == "经营数据"

    # 2.x 的内置样式只支持二维数组；格式不对的 sheets 报 INVALID_INPUT
    assert main(["create", "xlsx", "-o", "x.xlsx", "--data-file", "s.json", "--theme", "acks", "--json"]) == 1
    assert json.loads(capsys.readouterr().out)["error"]["code"] == "INVALID_INPUT"
    (tmp_path / "bad.json").write_text(json.dumps({"sheets": [{"rows": "x"}]}), encoding="utf-8")
    assert main(["create", "xlsx", "-o", "y.xlsx", "--data-file", "bad.json", "--theme", "slate", "--json"]) == 1
    assert json.loads(capsys.readouterr().out)["error"]["code"] == "INVALID_INPUT"


def test_acks_excel_unchanged_and_default_is_neutral(tmp_path):
    out = tmp_path / "acks.xlsx"
    result = create_excel("销售", DATA, str(out), theme="acks")
    assert result["theme"] == "acks" and result["rows"] == len(DATA)
    assert _sheet(out)["A1"].value == "区域"
    for theme in (None, "default"):  # 3.0 起默认主题是 neutral，default 是它的别名
        result = create_excel("销售", DATA, str(tmp_path / f"{theme}.xlsx"), theme=theme)
        assert result["theme"] == "neutral" and _sheet(tmp_path / f"{theme}.xlsx")["A1"].value == "销售"


def test_sheet_variants_are_validated(tmp_path):
    bad = tmp_path / "bad"
    bad.mkdir()
    (bad / "theme.json").write_text(json.dumps({"name": "bad", "extends": "slate"}), encoding="utf-8")
    (bad / "tokens.json").write_text(json.dumps({"layout": {"sheet": {"total": "box", "highlight": "nope"},
                                                            "doc": {"table": {"total": "x"}}}}), encoding="utf-8")
    report = validate_theme(str(bad), check_installed=False)
    errors = {(e["code"], e["path"]) for e in report["errors"]}
    assert ("BAD_VARIANT", "layout.sheet.total") in errors and ("BAD_VARIANT", "layout.doc.table.total") in errors
    assert ("BAD_COLOR_REF", "layout.sheet.highlight") in errors
