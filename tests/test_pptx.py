"""PPT：品牌参数、标题不重复、切换效果。"""

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

import acks_office
from acks_office.pptx import TRANSITIONS, add_transition_effects, create_pptx

SLIDES = [
    {"title": "经营回顾", "subtitle": "2026 Q3", "layout": "title"},
    {"title": "Quarterly Review", "layout": "title"},
    {"title": "核心结论", "subtitle": "营收与会员", "content": "营收增长 24%\n\n会员 12 万"},
    {"title": "空白页"},
]


def _texts(path):
    return [[shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]
            for slide in Presentation(str(path)).slides]


@pytest.fixture
def deck(tmp_path):
    def make(**kwargs):
        out = tmp_path / "deck.pptx"
        kwargs.setdefault("theme", "acks")  # 这些用例针对 2.x 的 ACKS 样式
        create_pptx("演示", SLIDES, str(out), **kwargs)
        return out
    return make


def test_acks_brand_is_kept(deck):
    flat = "\n".join(t for slide in _texts(deck()) for t in slide)
    assert "ACKS STUDIO" in flat and "CONFIDENTIAL" in flat


def test_default_theme_has_no_brand(deck):
    flat = "\n".join(t for slide in _texts(deck(theme=None)) for t in slide)
    assert "ACKS" not in flat and "CONFIDENTIAL" not in flat


def test_empty_brand_removes_brand_everywhere(deck):
    flat = "\n".join(t for slide in _texts(deck(brand_name="")) for t in slide)
    for word in ("ACKS", "爱驰科驶", "CONFIDENTIAL"):
        assert word not in flat


def test_custom_brand(deck):
    texts = _texts(deck(brand_name="栖木咖啡", footer_label="内部资料"))
    flat = "\n".join(t for slide in texts for t in slide)
    assert "栖木咖啡" in flat and "内部资料" in flat
    assert "ACKS" not in flat and "爱驰科驶" not in flat


def test_titles_appear_once_without_stray_spaces(deck):
    cover_cn, cover_en, content, empty = _texts(deck())
    assert "经营回顾" in cover_cn and "2026 Q3" in cover_cn
    assert "Quarterly Review" in cover_en
    assert sum("核心结论" in t for t in content) == 1 and "营收与会员" in content
    assert any("营收增长 24%" in t for t in content)
    assert sum("空白页" in t for t in empty) == 1


def _transitions(path):
    return [slide._element.findall(qn("p:transition")) for slide in Presentation(str(path)).slides]


def test_transitions_are_written_once_per_slide(deck, tmp_path):
    src, out = deck(), tmp_path / "fade.pptx"
    result = add_transition_effects(str(src), str(out), effect="push", duration=1.5)
    assert result == {"output_path": str(out), "slides": 4, "effect": "push"}
    add_transition_effects(str(out), str(out), effect="fade", duration=0.3)  # 再设一次会替换，不会叠加
    before = src.read_bytes()
    default = add_transition_effects(str(src), effect="wipe")  # 3.0：默认写到新文件，不改原文件
    assert default["output_path"] == str(src.with_name("deck_transitions.pptx")) and src.read_bytes() == before

    for found in _transitions(out):
        (transition,) = found
        assert transition.get("spd") == "fast"
        assert [child.tag for child in transition] == [qn("p:fade")]
        following = transition.getnext()
        assert following is None or following.tag in (qn("p:timing"), qn("p:extLst"))


@pytest.mark.parametrize("effect", TRANSITIONS)
def test_every_listed_effect_is_valid(deck, tmp_path, effect):
    out = tmp_path / f"{effect}.pptx"
    add_transition_effects(str(deck()), str(out), effect=effect)
    assert all(len(found) == 1 for found in _transitions(out))


def test_unknown_effect_is_rejected(deck):
    with pytest.raises(ValueError, match="不支持的切换效果"):
        add_transition_effects(str(deck()), effect="sparkle")


def test_extract_returns_text_per_slide(deck):
    assert "核心结论" in acks_office.extract(str(deck()))[3]
