"""xlsx.extract_data 回归测试：默认读取第一个工作表，始终返回行字典列表。"""

import datetime
import json
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from office_suite import OfficeSuite
from office_suite.xlsx import extract_data

DATA = Path(__file__).resolve().parent / "data"
FIRST = [{"name": "A", "value": 1}, {"name": "B", "value": 2}]
SECOND = [{"name": "C", "value": 3}]


@pytest.fixture
def workbook(tmp_path):
    """两个工作表：First 与 Second。"""
    path = tmp_path / "sample.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame(FIRST).to_excel(writer, sheet_name="First", index=False)
        pd.DataFrame(SECOND).to_excel(writer, sheet_name="Second", index=False)
    return str(path)


def test_defaults_to_first_sheet(workbook):
    assert extract_data(workbook) == FIRST


def test_explicit_none_reads_first_sheet(workbook):
    assert extract_data(workbook, sheet_name=None) == FIRST


@pytest.mark.parametrize("sheet", ["Second", 1])
def test_selects_sheet_by_name_or_index(workbook, sheet):
    assert extract_data(workbook, sheet_name=sheet) == SECOND


def test_forwards_read_excel_kwargs(workbook):
    assert extract_data(workbook, usecols=["value"]) == [{"value": 1}, {"value": 2}]


@pytest.mark.filterwarnings("error::pandas.errors.ParserWarning")
def test_converters_take_precedence_without_warning(workbook):
    assert extract_data(workbook, converters={"value": str}) == [
        {"name": "A", "value": "1"},
        {"name": "B", "value": "2"},
    ]


def test_unknown_sheet_name_raises(workbook):
    with pytest.raises(ValueError, match="Missing"):
        extract_data(workbook, sheet_name="Missing")


def test_rejects_multiple_sheets(workbook):
    with pytest.raises(TypeError, match="sheet_name"):
        extract_data(workbook, sheet_name=["First", "Second"])


def test_empty_sheet_returns_empty_list(tmp_path):
    path = tmp_path / "empty.xlsx"
    pd.DataFrame().to_excel(path, index=False)
    assert extract_data(str(path)) == []


def test_reads_legacy_xls():
    assert extract_data(str(DATA / "sample.xls")) == FIRST


def test_keeps_values_as_stored_in_excel(tmp_path):
    """文本型数字保持原样；空单元格不会把整列的 int / bool 变成 float。"""
    path = tmp_path / "ids.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "工号", "身份证号", "数量", "启用"])
    ws.append([1, "00123", "110101199003071234", 5, True])
    ws.append([2, None, None, None, None])
    wb.save(path)

    # 比较 JSON 文本才能区分 1 / 1.0 / True，以及 null / NaN
    assert json.dumps(extract_data(str(path)), ensure_ascii=False, allow_nan=False) == json.dumps([
        {"id": 1, "工号": "00123", "身份证号": "110101199003071234", "数量": 5, "启用": True},
        {"id": 2, "工号": None, "身份证号": None, "数量": None, "启用": None},
    ], ensure_ascii=False)


def test_dates_and_times_become_iso_strings(tmp_path):
    path = tmp_path / "dates.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["时间点", "时刻", "时长", datetime.datetime(2026, 1, 1)])
    ws.append([datetime.datetime(2026, 9, 1, 8, 30), datetime.time(9, 30), datetime.timedelta(hours=26), 10])
    ws.append([None, None, None, 20])
    ws["C2"].number_format = "[h]:mm:ss"
    wb.save(path)

    rows = extract_data(str(path))

    assert rows == [
        {"时间点": "2026-09-01T08:30:00", "时刻": "09:30:00", "时长": "P1DT2H0M0S", "2026-01-01T00:00:00": 10},
        {"时间点": None, "时刻": None, "时长": None, "2026-01-01T00:00:00": 20},
    ]
    json.dumps(rows, allow_nan=False)


def test_suite_default_call_succeeds_and_leaves_source_untouched(workbook):
    with open(workbook, "rb") as f:
        before = f.read()

    assert OfficeSuite().extract_data(workbook) == {"success": True, "data": FIRST}

    with open(workbook, "rb") as f:
        assert f.read() == before
