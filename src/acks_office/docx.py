import io
import os
from copy import deepcopy
from typing import Optional, List, Dict, Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table as DocxTable



def create_word(title: Optional[str], content: str, output_path: str,
                theme: Any = None,
                brand_name: Optional[str] = None,
                footer_label: Optional[str] = None,
                **kwargs) -> Dict[str, Any]:
    """
    创建 Word 文档。

    Args:
        title: 文档标题；为空时用正文 front matter 里的 title
        content: 文档内容，Markdown：标题、粗体/斜体、链接、列表与任务清单、表格、引文与提示块、代码块、图片；
            开头可以写 front matter（kicker、subtitle、author、date、version、issue、lede 等）
        output_path: 输出文件路径
        theme: 主题名称（neutral、slate、folio 或已安装的主题）、主题目录或 Theme 对象；
            "default" 是 neutral 的别名
        brand_name: 品牌名，传 "" 去品牌化
        footer_label: 页脚左侧文字，默认按主题模板生成
        subtitle 等元数据: 见 front matter
        base_dir: 图片相对路径的基准目录，默认当前目录
        font_policy: "local"（缺主题字体时改用本机字体，默认）或 "theme"（总是写主题字体名）
    """
    from .render.common import META_KEYS
    from .render.word import render_word
    meta = {key: kwargs[key] for key in kwargs if key in META_KEYS or key not in _RENDER_OPTIONS}
    if brand_name is not None:
        meta["brand"] = brand_name
    return render_word(title, content, output_path, theme=theme, base_dir=kwargs.get("base_dir"),
                       font_policy=kwargs.get("font_policy", "local"), footer_label=footer_label, **meta)


_RENDER_OPTIONS = ("base_dir", "font_policy", "theme", "create_chart", "font", "bold_font")


# ---------------------------------------------------------------- 水印、提取、合并

def add_watermark(input_path: str, watermark_text: str, output_path: Optional[str] = None,
                  font_size: int = 48, color: tuple = (200, 200, 200),
                  italic: bool = True, **kwargs) -> Dict[str, Any]:
    """
    给 Word 文档添加水印（页眉斜体灰色大字，简化版）。

    Args:
        input_path: 输入文件路径
        watermark_text: 水印文本
        output_path: 输出文件路径，默认为 <原名>_watermarked.docx，不覆盖原文件
        font_size: 字号
        color: RGB 颜色元组
        italic: 是否斜体
    """
    if not output_path:
        from .api import watermarked_path
        output_path = watermarked_path(input_path)

    doc = Document(input_path)

    section = doc.sections[0]
    header = section.header

    watermark_para = header.add_paragraph()
    watermark_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = watermark_para.add_run(watermark_text)
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color)
    run.bold = True
    run.italic = italic

    doc.save(output_path)

    return {"output_path": output_path}


def extract_text(input_path: str, **kwargs) -> str:
    """
    提取 Word 文档中的文本：段落与表格按原文顺序输出，表格每行一行、单元格用 " | " 分隔
    """
    doc = Document(input_path)
    lines: List[str] = []
    for item in doc.iter_inner_content():
        if isinstance(item, DocxTable):
            lines.extend(_table_lines(item))
        else:
            lines.append(item.text)
    return '\n'.join(lines)


def _table_lines(table: DocxTable) -> List[str]:
    lines = []
    for row in table.rows:
        seen, cells = [], []
        for cell in row.cells:
            if any(cell._tc is tc for tc in seen):  # 合并单元格只取一次
                continue
            seen.append(cell._tc)
            cells.append(cell.text.strip().replace("\n", " "))
        lines.append(" | ".join(cells))
    return lines


_REL_ATTRS = (qn("r:embed"), qn("r:link"), qn("r:id"))


def _copy_relationships(element, src_part, dst_part) -> List[Dict]:
    """复制元素引用的图片与外部链接关系，并把元素里的关系 ID 改成目标文档里的新 ID。"""
    warnings, mapping = [], {}
    for el in element.iter():
        for attr in _REL_ATTRS:
            rid = el.get(attr)
            if not rid or rid not in src_part.rels:
                continue
            if rid not in mapping:
                rel = src_part.rels[rid]
                if rel.is_external:
                    mapping[rid] = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                elif rel.reltype == RT.IMAGE:
                    mapping[rid], _ = dst_part.get_or_add_image(io.BytesIO(rel.target_part.blob))
                else:
                    mapping[rid] = None
                    warnings.append({"code": "MERGE_PARTIAL",
                                     "message": f"未能合并内嵌对象：{rel.reltype.rsplit('/', 1)[-1]}"})
            if mapping[rid]:
                el.set(attr, mapping[rid])
    if element.find(".//" + qn("w:numPr")) is not None:
        warnings.append({"code": "MERGE_NUMBERING", "message": "列表编号沿用第一份文档的编号定义，格式可能与原文档不同"})
    return warnings


def merge_documents(input_paths: List[str], output_path: str, **kwargs) -> Dict[str, Any]:
    """
    合并多个 Word 文档：以第一份为底（沿用其样式与页面设置），后续文档另起一页接在后面。
    任一文件不存在时报错，不会只合并一部分。
    """
    missing = [p for p in input_paths if not os.path.exists(p)]
    if missing:
        raise FileNotFoundError(f"要合并的文件不存在：{', '.join(missing)}")
    merged = None
    merged_count = 0
    warnings: List[Dict] = []

    for path in input_paths:
        if merged is None:
            merged = Document(path)
        else:
            source = Document(path)
            body = merged.element.body
            anchor = body.sectPr
            page_break = OxmlElement("w:p")
            run = OxmlElement("w:r")
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            run.append(br)
            page_break.append(run)
            for element in [page_break] + [e for e in source.element.body if e.tag != qn("w:sectPr")]:
                if element is not page_break:
                    element = deepcopy(element)
                    warnings.extend(_copy_relationships(element, source.part, merged.part))
                if anchor is not None:
                    anchor.addprevious(element)
                else:
                    body.append(element)
        merged_count += 1

    if merged is None:
        merged = Document()
    merged.save(output_path)

    result: Dict[str, Any] = {"output_path": output_path, "merged_count": merged_count}
    if warnings:
        result["warnings"] = [dict(t) for t in {tuple(w.items()) for w in warnings}]
    return result
