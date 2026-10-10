from pathlib import Path
from typing import Optional, List, Dict, Any
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

TRANSITIONS = ("fade", "push", "wipe", "split", "cover", "pull", "dissolve", "cut", "zoom", "random")


def create_pptx(title: str, slides: List[Dict[str, Any]], output_path: str,
                theme: Any = None,
                brand_name: Optional[str] = None,
                **kwargs) -> Dict[str, Any]:
    """
    创建 PowerPoint 演示文稿。

    Args:
        title: 演示文稿标题
        slides: 幻灯片列表。每页一个字典，layout 为 title（封面）、section（章节页）、content（内容页，
            content 或 bullets）、data（关键数字 kpis + 条形图 chart）、table（表格）、number（大数字）、
            quote（引文）、image（图片），其余字段见命令参考
        output_path: 输出路径
        theme: 主题名称（neutral、slate、folio 或已安装的主题）、主题目录或 Theme 对象；
            "default" 是 neutral 的别名
        brand_name: 品牌名，传 "" 去品牌化
        footer_label: 页脚左侧文字，默认按主题模板生成
        其余关键字参数（kicker、classification、publication、issue、season、date、author 等）作为元数据
    """
    from .render.slides import render_slides
    options = ("base_dir", "font_policy", "footer_label")
    meta = {k: v for k, v in kwargs.items() if k not in options}
    if brand_name is not None:
        meta["brand"] = brand_name
    return render_slides(title, slides, output_path, theme=theme, base_dir=kwargs.get("base_dir"),
                         font_policy=kwargs.get("font_policy", "local"), footer_label=kwargs.get("footer_label"),
                         **meta)


def add_transition_effects(input_path: str, output_path: Optional[str] = None, effect: str = "fade",
                           duration: float = 0.5, **kwargs) -> Dict[str, Any]:
    """
    给每张幻灯片添加切换效果
    Args:
        effect: fade / push / wipe / split / cover / pull / dissolve / cut / zoom / random
        duration: 秒，映射为 PowerPoint 的快（≤0.5）/ 中（≤1）/ 慢 三档
        output_path: 输出路径，默认为 <原名>_transitions.pptx，不覆盖原文件（要改原文件时显式传入原路径）
    """
    if effect not in TRANSITIONS:
        raise ValueError(f"不支持的切换效果：{effect}；可选：{', '.join(TRANSITIONS)}")
    if not output_path:
        source = Path(input_path)
        output_path = str(source.with_name(f"{source.stem}_transitions{source.suffix}"))

    speed = "fast" if duration <= 0.5 else "med" if duration <= 1.0 else "slow"
    prs = Presentation(input_path)
    for slide in prs.slides:
        sld = slide._element
        for old in sld.findall(qn("p:transition")):
            sld.remove(old)
        transition = parse_xml(f'<p:transition {nsdecls("p")} spd="{speed}"><p:{effect}/></p:transition>')
        # 元素顺序：cSld、clrMapOvr、transition、timing
        timing = sld.find(qn("p:timing"))
        if timing is not None:
            timing.addprevious(transition)
        else:
            anchor = sld.find(qn("p:clrMapOvr"))
            (anchor if anchor is not None else sld.find(qn("p:cSld"))).addnext(transition)

    prs.save(output_path)
    return {"output_path": output_path, "slides": len(prs.slides), "effect": effect}


def extract_text(input_path: str, **kwargs) -> Dict[int, str]:
    """
    提取PPT中的文本，返回字典，key是幻灯片编号，value是该页文本
    """
    prs = Presentation(input_path)
    result = {}

    for i, slide in enumerate(prs.slides, 1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text.append(shape.text)
        result[i] = '\n'.join(slide_text)

    return result
