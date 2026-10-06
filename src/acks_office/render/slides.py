"""主题驱动的 PPT 渲染（16:9，960 × 540 磅）。

幻灯片版式（layout）：title 封面、section 章节页、content 内容页、data 数据页（关键数字 + 条形图）、
table 表格页、number 大数字页、quote 引文页、image 图片页。封面与章节页的样子由主题变体决定：
封面 simple / split / issue，章节页 light / dark，内容页 stacked / columns。
"""

import os
from typing import Any, Dict, List, Optional, Sequence, Tuple

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from ..themes import Theme, load_theme
from . import measure

W, H = 960.0, 540.0
ROLES = ("body", "title", "heading", "label", "number", "display", "quote", "code")
LAYOUTS = ("title", "section", "content", "data", "table", "number", "quote", "image")
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # PowerPoint 内置「无样式，无网格」
Run = Tuple[str, Dict[str, Any]]


def emu(points: float) -> Emu:
    return Emu(int(round(points * 12700)))


def para(*runs: Run, align: str = "left", leading: Optional[float] = None, before: float = 0,
         after: float = 0, indent: float = 0) -> Dict[str, Any]:
    return {"runs": [r for r in runs if r[0]], "align": align, "leading": leading, "before": before,
            "after": after, "indent": indent}


def r(text: str, **style: Any) -> Run:
    return (text, style)


class SlideRenderer:
    def __init__(self, theme: Theme, meta: Dict[str, Any], base_dir: str, policy: str = "local",
                 footer_label: Optional[str] = None):
        self.t = theme
        self.meta = meta
        self.values = theme.chrome_values(meta)
        self.base_dir = base_dir
        self.footer_label = footer_label
        self.warnings: List[Dict[str, str]] = []
        self.fonts = {role: theme.font(role, policy) for role in ROLES}
        for choice in self.fonts.values():
            for note in choice.substitutions:
                if not any(w.get("family") == note["family"] for w in self.warnings):
                    self.warnings.append(note)
        self.layout = theme.get("layout.slide")
        margin = self.layout["margin"]
        self.mt, self.mr, self.mb, self.ml = margin["top"], margin["right"], margin["bottom"], margin["left"]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = emu(W), emu(H)
        self.sections = 0

    # ------------------------------------------------------------ 基本元素

    def sz(self, key: str) -> float:
        return float(self.t.get(f"type.slide.{key}"))

    def lead(self, key: str) -> float:
        return float(self.t.get(f"leading.{key}"))

    def rgb(self, color: str) -> RGBColor:
        return RGBColor.from_string(self.t.hex(color))

    def new_slide(self, background: str = "slide_paper"):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.rgb(background)
        return slide

    def rect(self, slide, x: float, y: float, w: float, h: float, color: str):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def hline(self, slide, x: float, y: float, w: float, color: str = "rule", weight: float = 0.75):
        return self.rect(slide, x, y, w, weight, color)

    def style_run(self, run, role: str = "body", size: float = 12, color: str = "ink",
                  bold: Optional[bool] = None, italic: bool = False, track: float = 0) -> None:
        choice = self.fonts[role]
        font = run.font
        font.size = Pt(size)
        font.bold = choice.bold if bold is None else bold
        font.italic = italic
        font.color.rgb = self.rgb(color)
        font.name = choice.en
        rpr = run._r.get_or_add_rPr()
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(rpr, qn("a:ea"))
            rpr.find(qn("a:latin")).addnext(ea)
        ea.set("typeface", choice.cn)
        if track:
            rpr.set("spc", str(int(round(track * size * 100))))

    def box(self, slide, x: float, y: float, w: float, h: float, paragraphs: Sequence[Dict[str, Any]],
            anchor: str = "top", wrap: bool = True):
        shape = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(max(h, 1)))
        tf = shape.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        first = True
        for spec in paragraphs:
            if not spec["runs"]:
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[spec["align"]]
            size = max(style.get("size", 12) for _, style in spec["runs"])
            if spec["leading"]:
                p.line_spacing = Pt(size * spec["leading"])
            if spec["before"]:
                p.space_before = Pt(spec["before"])
            if spec["after"]:
                p.space_after = Pt(spec["after"])
            if spec["indent"]:
                p._pPr.set("marL", str(int(emu(spec["indent"]))))
                p._pPr.set("indent", str(-int(emu(spec["indent"]))))
            for text, style in spec["runs"]:
                run = p.add_run()
                run.text = text
                self.style_run(run, **style)
        return shape

    def height(self, text: str, size: float, leading: float, width: float, role: str,
               bold: Optional[bool] = None) -> float:
        lines = measure.wrap_count(text, size, width, self.fonts[role], bold) if text else 0
        return lines * size * leading

    def fit(self, text: str, size: float, width: float, role: str, lines: int = 2,
            minimum: Optional[float] = None) -> float:
        return measure.fit_size(text, size, width, self.fonts[role], lines, minimum)

    # ------------------------------------------------------------ 页脚与通用标题

    def footer(self, slide, number: int, color: str = "muted", rule: str = "rule") -> None:
        if not self.layout["footer"]:
            return
        size = self.sz("footer")
        y = H - self.mb - size * 1.3
        self.hline(slide, self.ml, y - 10.5, W - self.ml - self.mr, rule)
        page = f"{number:02d}"
        left = ([("text", self.footer_label)] if self.footer_label is not None
                else self.t.template("slide.footer.left", self.values))
        right = self.t.template("slide.footer.right", self.values)
        left_text = "".join(v if k == "text" else page for k, v in left)
        right_text = "".join(v if k == "text" else page for k, v in right)
        width = W - self.ml - self.mr
        self.box(slide, self.ml, y, width * 0.75, size * 1.4, [para(r(left_text, role="label", size=size,
                                                                      color=color, bold=False))])
        self.box(slide, self.ml + width * 0.75, y, width * 0.25, size * 1.4,
                 [para(r(right_text, role="number", size=size, color=color, bold=False), align="right")])

    def heading(self, slide, spec: Dict[str, Any], width: Optional[float] = None) -> float:
        """标题与一句结论（副标题），返回下方内容的起始 y。"""
        width = width or (W - self.ml - self.mr)
        title = str(spec.get("title", ""))
        size = self.fit(title, self.sz("title"), width, "title", lines=2, minimum=self.sz("title") * 0.7)
        leading = self.lead("slide_title")
        height = self.height(title, size, leading, width, "title")
        self.box(slide, self.ml, self.mt, width, height + 4,
                 [para(r(title, role="title", size=size), leading=leading)])
        y = self.mt + height
        subtitle = str(spec.get("subtitle", ""))
        if subtitle:
            sub = self.sz("subtitle")
            sub_height = self.height(subtitle, sub, 1.5, width, "body")
            self.box(slide, self.ml, y + 9, width, sub_height + 4,
                     [para(r(subtitle, role="body", size=sub, color="muted"), leading=1.5)])
            y += 9 + sub_height
        return y

    def label_row(self, slide, label: str) -> float:
        """folio 的页眉标签：「营收 · REVENUE」，下方一条细线；返回下方内容起始 y。"""
        head, _, tail = label.partition(" · ")
        size = self.sz("big_label")
        runs = [r(head, role="label", size=size, track=0.08)]
        if tail:
            runs += [r("  ·  ", role="label", size=size * 0.87, color="muted", bold=False),
                     r(tail.upper(), role="display", size=size * 0.8, color="muted", bold=True,
                       track=self.t.get("tracking.caps"))]
        self.box(slide, self.ml, self.mt, W - self.ml - self.mr, size * 1.5, [para(*runs)])
        y = self.mt + size * 1.5 + 10.5
        self.hline(slide, self.ml, y, W - self.ml - self.mr)
        return y

    # ------------------------------------------------------------ 封面

    def cover(self, spec: Dict[str, Any], number: int) -> None:
        variant = self.layout["cover"]
        {"split": self.cover_split, "issue": self.cover_issue}.get(variant, self.cover_simple)(spec, number)

    def _brand_row(self, slide, x: float, width: float, y: float) -> None:
        brand, classification = self.values.get("brand", ""), self.values.get("classification", "")
        if brand:
            self.box(slide, x, y, width * 0.7, 24, [para(r(brand, role="heading", size=self.sz("brand"),
                                                          track=0.02))])
        if classification:
            self.box(slide, x + width * 0.5, y + 4, width * 0.5, 20,
                     [para(r(classification, role="label", size=self.sz("classification"), color="muted",
                             bold=False), align="right")])

    def _title_stack(self, slide, spec: Dict[str, Any], x: float, width: float, top: float, bottom: float,
                     center: bool = True) -> None:
        """强调短线、类型、大标题、副标题；在 top 与 bottom 之间居中（或贴底）。"""
        kicker, title = str(spec.get("kicker") or self.meta.get("kicker") or ""), str(spec.get("title", ""))
        subtitle = str(spec.get("subtitle", ""))
        title_size = self.fit(title, self.sz("cover_title"), width, "title", lines=2,
                              minimum=self.sz("cover_title") * 0.5)
        lead = self.lead("slide_cover_title")
        title_h = self.height(title, title_size, lead, width, "title")
        kicker_h = self.sz("cover_kicker") * 1.3 if kicker else 0
        sub_h = self.height(subtitle, self.sz("cover_subtitle"), 1.5, width, "body")
        gap = 15
        total = 3 + gap + (kicker_h + gap if kicker else 0) + title_h + (gap + sub_h if subtitle else 0)
        y = top + max(0, (bottom - top - total) / 2) if center else bottom - total
        self.rect(slide, x, y, 42, 3, "accent")
        y += 3 + gap
        if kicker:
            self.box(slide, x, y, width, kicker_h + 2, [para(r(kicker, role="label", size=self.sz("cover_kicker"),
                                                               color="accent"))])
            y += kicker_h + gap
        self.box(slide, x, y, width, title_h + 6, [para(r(title, role="title", size=title_size), leading=lead)])
        y += title_h + gap
        if subtitle:
            self.box(slide, x, y, width, sub_h + 4, [para(r(subtitle, role="body", size=self.sz("cover_subtitle"),
                                                            color="muted"), leading=1.5)])

    def _cover_footer_line(self, slide, x: float, width: float, text: str) -> float:
        size = self.sz("cover_footer")
        y = H - 48 - size * 1.3
        self.hline(slide, x, y - 15, width)
        if text:
            self.box(slide, x, y, width, size * 1.4, [para(r(text, role="label", size=size, color="muted",
                                                             bold=False))])
        return y - 15

    def _byline(self, spec: Dict[str, Any]) -> str:
        parts = [spec.get("footer")] if spec.get("footer") else [self.meta.get(k) for k in ("author", "date")]
        return " · ".join(str(p) for p in parts if p)

    def cover_simple(self, spec: Dict[str, Any], number: int) -> None:
        slide = self.new_slide()
        width = W - self.ml - self.mr
        self._brand_row(slide, self.ml, width, 54)
        rule_y = self._cover_footer_line(slide, self.ml, width, self._byline(spec))
        self._title_stack(slide, spec, self.ml, width * 0.86, 90, rule_y - 20)
        self.notes(slide, spec)

    def cover_split(self, spec: Dict[str, Any], number: int) -> None:
        """slate：左侧标题，右侧浅底竖栏放关键数字（没有 kpis 时标题占满整页）。"""
        slide = self.new_slide()
        kpis = spec.get("kpis") or []
        panel = 330 if kpis else 0
        left, right_edge = self.ml, W - panel - (48 if kpis else self.mr)
        width = right_edge - left
        self._brand_row(slide, left, width, 54)
        rule_y = self._cover_footer_line(slide, left, width, self._byline(spec))
        self._title_stack(slide, spec, left, width, 90, rule_y - 20)
        if kpis:
            self.rect(slide, W - panel, 0, panel, H, "surface")
            self._kpi_stack(slide, kpis, W - panel + 48, panel - 96, 54, H - 54, big=True)
        self.notes(slide, spec)

    def _kpi_stack(self, slide, kpis: List[Dict[str, Any]], x: float, width: float, top: float, bottom: float,
                   big: bool = False) -> None:
        label_size = self.sz("cover_kpi_label" if big else "kpi_label")
        value_size = self.sz("cover_kpi_value" if big else "kpi_value")
        unit_size = self.sz("cover_kpi_unit" if big else "kpi_unit")
        delta_size = self.sz("kpi_delta")
        pad = 16.5 if big else 10.5
        heights = []
        for kpi in kpis:
            h = pad * 2 + label_size * 1.3 + 7 + value_size * 1.08
            if kpi.get("delta"):
                h += 4 + delta_size * 1.3
            heights.append(h)
        total = sum(heights)
        y = top + max(0, (bottom - top - total) / 2) if big else top
        if not big:
            self.hline(slide, x, y, width)
        for i, (kpi, h) in enumerate(zip(kpis, heights, strict=True)):
            cy = y + pad
            self.box(slide, x, cy, width, label_size * 1.4,
                     [para(r(str(kpi.get("label", "")), role="label", size=label_size, color="muted", bold=False))])
            cy += label_size * 1.3 + 7
            color = "accent" if kpi.get("highlight") else "ink"
            runs = [r(str(kpi.get("value", "")), role="number", size=value_size, color=color)]
            if kpi.get("unit"):
                runs.append(r(" " + str(kpi["unit"]), role="label", size=unit_size, color=color))
            self.box(slide, x, cy, width, value_size * 1.15, [para(*runs, leading=1.0)])
            cy += value_size * 1.08
            if kpi.get("delta"):
                self.box(slide, x, cy + 4, width, delta_size * 1.4,
                         [para(r(str(kpi["delta"]), role="label", size=delta_size, color="accent", bold=False))])
            y += h
            if not big or i < len(kpis) - 1:
                self.hline(slide, x, y, width)

    def cover_issue(self, spec: Dict[str, Any], number: int) -> None:
        """folio：刊名与期号、左侧竖线与大标题、右侧超大期号，底部一条细线。"""
        slide = self.new_slide()
        x, width = self.ml, W - self.ml - self.mr
        top = 45.0
        publication = str(self.values.get("publication") or self.values.get("brand", ""))
        issue, season = str(self.meta.get("issue", "")), str(self.meta.get("season", ""))
        label = " — ".join(p for p in ((f"ISSUE {issue}" if issue else ""), season.upper()) if p)
        runs = [r(publication, role="label", size=self.sz("brand"), track=0.08)]
        if label:
            runs += [r("　　", role="label", size=self.sz("brand")),
                     r(label, role="display", size=self.sz("issue_label"), color="muted", bold=True,
                       track=self.t.get("tracking.caps") * 0.9)]
        self.box(slide, x, top, width * 0.75, 20, [para(*runs)])
        if self.meta.get("date"):
            self.box(slide, x + width * 0.6, top, width * 0.4, 20,
                     [para(r(str(self.meta["date"]), role="display", size=self.sz("brand"), color="muted",
                             bold=False), align="right")])
        size = self.sz("footer")
        foot_y = H - 30 - size * 1.3
        self.hline(slide, x, foot_y - 10.5, width)
        note = str(spec.get("footer") or self.meta.get("kicker") or "")
        if note:
            self.box(slide, x, foot_y, width * 0.6, size * 1.4,
                     [para(r(note, role="label", size=size, color="muted", bold=False))])
        tag = str(self.values.get("brand") or "").upper()
        if tag:
            self.box(slide, x + width * 0.4, foot_y, width * 0.6, size * 1.4,
                     [para(r(tag, role="display", size=size * 0.86, color="muted", bold=True,
                             track=self.t.get("tracking.caps") * 1.2), align="right")])
        bottom = foot_y - 10.5 - 27
        number_w = 300.0 if issue else 0
        if issue:
            big = self.sz("issue_number")
            self.box(slide, W - self.mr - number_w, bottom - big * 0.86, number_w, big * 0.9,
                     [para(r(issue, role="display", size=big, color="accent", bold=False), align="right",
                           leading=0.86)], anchor="bottom")
            no = self.sz("issue_no")
            self.box(slide, W - self.mr - number_w, bottom - big * 0.86 - no * 1.4, number_w, no * 1.4,
                     [para(r(self.t.label("issue"), role="display", size=no, color="muted", bold=False,
                             italic=True), align="right")])
        title = str(spec.get("title", ""))
        text_w = width - number_w - 28 - 2.25
        title_size = self.fit(title, self.sz("cover_title"), text_w, "title", lines=2,
                              minimum=self.sz("cover_title") * 0.5)
        lead = self.lead("slide_cover_title")
        title_h = self.height(title, title_size, lead, text_w, "title")
        subtitle = str(spec.get("subtitle", "") or self.meta.get("lede", ""))
        sub_size = self.sz("cover_subtitle")
        sub_w = min(text_w, 435)
        sub_h = self.height(subtitle, sub_size, 1.7, sub_w, "body") if subtitle else 0
        block = title_h + (21 + sub_h if subtitle else 0)
        y = bottom - block
        self.rect(slide, x, y, 2.25, block, "accent")
        tx = x + 2.25 + 21
        self.box(slide, tx, y, text_w, title_h + 6, [para(r(title, role="title", size=title_size), leading=lead)])
        if subtitle:
            self.box(slide, tx, y + title_h + 21, sub_w, sub_h + 4,
                     [para(r(subtitle, role="body", size=sub_size, color="muted"), leading=1.7)])
        self.notes(slide, spec)

    # ------------------------------------------------------------ 章节页

    def section(self, spec: Dict[str, Any], number: int) -> None:
        self.sections += 1
        dark = self.layout["section"] == "dark"
        slide = self.new_slide("dark" if dark else "slide_paper")
        ink, soft, muted = ("on_dark", "on_dark_soft", "on_dark_muted") if dark else ("ink", "muted", "muted")
        accent = "accent_on_dark" if dark else "accent"
        x, width = self.ml, W - self.ml - self.mr
        kicker = str(spec.get("kicker") or f"第 {self.sections} 部分")
        head, _, tail = kicker.partition(" · ")
        kicker_size = self.sz("section_kicker")
        runs = [r(head, role="display" if tail else "label", size=kicker_size, color=muted, bold=True,
                  track=self.t.get("tracking.caps") if tail else 0.08)]
        if tail:
            runs += [r("  ·  ", role="label", size=kicker_size, color=muted, bold=False),
                     r(tail, role="label", size=kicker_size * 1.08, color=muted, track=0.08)]
        self.box(slide, x, 48, width, 20, [para(*runs)])
        bottom = H - 60
        label = str(spec.get("number") or f"{self.sections:02d}")
        big = self.sz("section_number")
        number_w = measure.text_width(label, big, self.fonts["display"], False) + 12
        self.box(slide, x, bottom - big * 0.86, number_w, big * 0.9,
                 [para(r(label, role="display", size=big, color=accent, bold=False), leading=0.86)],
                 anchor="bottom")
        tx = x + number_w + 36
        tw = W - self.mr - tx
        title, title_en = str(spec.get("title", "")), str(spec.get("title_en", ""))
        size = self.fit(title, self.sz("section_title"), tw * 0.75 if title_en else tw, "title", lines=1,
                        minimum=self.sz("section_title") * 0.45)
        subtitle = str(spec.get("subtitle", ""))
        sub_size = self.sz("section_sub")
        sub_h = self.height(subtitle, sub_size, 1.7, tw, "body") if subtitle else 0
        y = bottom - 4 - sub_h - (14 if subtitle else 0) - size * 1.05
        runs = [r(title, role="title", size=size, color=ink, track=0.04)]
        if title_en:
            runs += [r("  " + title_en, role="display", size=self.sz("section_title_en"), color=muted, bold=False,
                       italic=True)]
        self.box(slide, tx, y, tw, size * 1.1, [para(*runs, leading=1.05)])
        if subtitle:
            self.box(slide, tx, y + size * 1.05 + 14, tw, sub_h + 4,
                     [para(r(subtitle, role="body", size=sub_size, color=soft), leading=1.7)])
        self.notes(slide, spec)

    # ------------------------------------------------------------ 内容页

    def _lines(self, spec: Dict[str, Any]) -> List[str]:
        if spec.get("bullets"):
            return [str(b) for b in spec["bullets"]]
        text = str(spec.get("content", "") or "")
        return [line.strip().lstrip("•-*·").strip() for line in text.split("\n") if line.strip()]

    def _bullets(self, slide, lines: List[str], x: float, y: float, width: float, bottom: float,
                 role: str = "body", marker: bool = True) -> None:
        size = self.sz("body")
        lead = self.lead("slide_body")
        available = bottom - y
        while size > 10:
            total = sum(self.height(line, size, lead, width - (size if marker else 0), role) + size * 0.5
                        for line in lines)
            if total <= available:
                break
            size *= 0.93
        paragraphs = []
        for line in lines:
            runs = [r("• ", role="body", size=size, color="accent", bold=False)] if marker else []
            runs.append(r(line, role=role, size=size, bold=False))
            paragraphs.append(para(*runs, leading=lead, after=size * 0.5, indent=size * 0.9 if marker else 0))
        self.box(slide, x, y, width, bottom - y, paragraphs)

    def content(self, spec: Dict[str, Any], number: int) -> None:
        slide = self.new_slide()
        width = W - self.ml - self.mr
        bottom = H - self.mb - self.sz("footer") * 1.3 - 24
        lines = self._lines(spec)
        if self.layout["content"] == "columns":
            label = str(spec.get("label") or spec.get("kicker") or self.meta.get("short_title") or "")
            y = self.label_row(slide, label) + 30 if label else self.mt
            left_w = 255.0
            title = str(spec.get("title", ""))
            size = self.fit(title, self.sz("title"), left_w, "title", lines=3, minimum=self.sz("title") * 0.6)
            self.box(slide, self.ml, y, left_w, bottom - y,
                     [para(r(title, role="title", size=size), leading=self.lead("slide_title"))])
            rx = self.ml + left_w + 48
            rw = W - self.mr - rx
            if spec.get("subtitle"):
                sub = str(spec["subtitle"])
                sub_h = self.height(sub, self.sz("subtitle"), 1.7, rw, "body")
                self.box(slide, rx, y, rw, sub_h + 4,
                         [para(r(sub, role="body", size=self.sz("subtitle"), color="muted"), leading=1.7)])
                y += sub_h + 18
            self._bullets(slide, lines, rx, y, rw, bottom, marker=bool(spec.get("bullets")))
        else:
            y = self.heading(slide, spec) + 33
            self._bullets(slide, lines, self.ml, y, width, bottom)
        self.footer(slide, number)
        self.notes(slide, spec)

    def _bars(self, slide, chart: Dict[str, Any], x: float, y: float, width: float, bottom: float) -> None:
        """横向条形图（形状绘制）：类别、条、数值；highlight 的那一项用强调色。"""
        data = [(str(row[0]), float(row[1])) for row in chart.get("data", [])]
        if not data:
            return
        title, unit = str(chart.get("title", "")), str(chart.get("unit", ""))
        head = self.sz("chart_title")
        if title or unit:
            self.box(slide, x, y, width * 0.7, head * 1.4, [para(r(title, role="label", size=head))])
            if unit:
                self.box(slide, x + width * 0.5, y + 2, width * 0.5, head * 1.4,
                         [para(r(unit, role="label", size=self.sz("small"), color="muted", bold=False),
                               align="right")])
            y += head * 1.4 + 9
            self.hline(slide, x, y, width)
            y += 19.5
        highlight = chart.get("highlight", 0)
        label_w, value_w, gap = 42.0, 48.0, 12.0
        note = str(chart.get("note", ""))
        note_h = self.sz("chart_note") * 1.5 if note else 0
        bar_h = 22.5
        step = min(bar_h + 10.5, max(bar_h + 4, (bottom - note_h - y) / max(len(data), 1)))
        peak = max(v for _, v in data) or 1
        for i, (name, value) in enumerate(data):
            on = highlight == i or highlight == name
            cy = y + i * step
            self.box(slide, x, cy, label_w, bar_h, [para(r(name, role="label", size=self.sz("chart_label"),
                                                           bold=False))], anchor="middle")
            bar_x = x + label_w + gap
            bar_w = (width - label_w - value_w - gap * 2) * max(value, 0) / peak
            if bar_w > 0:
                self.rect(slide, bar_x, cy, bar_w, bar_h, "accent" if on else "chart_muted")
            text = chart.get("format", "{:.1f}").format(value)
            self.box(slide, x + width - value_w, cy, value_w, bar_h,
                     [para(r(text, role="number", size=self.sz("chart_value"), color="accent" if on else "ink"),
                           align="right")], anchor="middle")
        if note:
            self.box(slide, x, bottom - note_h, width, note_h,
                     [para(r(note, role="label", size=self.sz("chart_note"), color="muted", bold=False))],
                     anchor="bottom")

    def data(self, spec: Dict[str, Any], number: int) -> None:
        slide = self.new_slide()
        y = self.heading(slide, spec) + 27
        bottom = H - self.mb - self.sz("footer") * 1.3 - 24
        kpis = spec.get("kpis") or []
        x = self.ml
        if kpis:
            self._kpi_stack(slide, kpis, x, 255, y, bottom)
            x += 255 + 60
        if spec.get("chart"):
            self._bars(slide, spec["chart"], x, y, W - self.mr - x, bottom)
        elif self._lines(spec):
            self._bullets(slide, self._lines(spec), x, y, W - self.mr - x, bottom)
        self.footer(slide, number)
        self.notes(slide, spec)

    # ------------------------------------------------------------ 表格页

    def table(self, spec: Dict[str, Any], number: int) -> None:
        slide = self.new_slide()
        y = self.heading(slide, spec) + 33
        bottom = H - self.mb - self.sz("footer") * 1.3 - 24
        table = spec.get("table") or {}
        header = [str(h) for h in table.get("header", [])]
        rows = [[str(c) for c in row] for row in table.get("rows", [])]
        cols = max([len(header)] + [len(row) for row in rows]) if (header or rows) else 0
        if cols:
            width = W - self.ml - self.mr
            weights = table.get("widths") or [max([measure.text_width(t, 12, self.fonts["body"], False)
                                                   for t in [header[i] if i < len(header) else ""]
                                                   + [row[i] if i < len(row) else "" for row in rows]] + [24])
                                              for i in range(cols)]
            total = sum(weights)
            widths = [width * w / total for w in weights]
            count = len(rows) + (1 if header else 0)
            shape = slide.shapes.add_table(count, cols, emu(self.ml), emu(y), emu(width), emu(min(bottom - y, 40 * count)))
            tbl = shape.table
            style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
            if style is None:
                style = etree.SubElement(tbl._tbl.tblPr, qn("a:tableStyleId"))
            style.text = NO_STYLE_TABLE
            tbl.first_row = False
            tbl.horz_banding = False
            for i, w in enumerate(widths):
                tbl.columns[i].width = emu(w)
            numbered = bool(rows) and all(row and row[0].strip().isdigit() for row in rows)
            head_fill = self.t.get("layout.doc.table.head") == "fill"
            all_rows = ([("head", header)] if header else []) + [("row", row) for row in rows]
            for ri, (kind, cells) in enumerate(all_rows):
                for ci in range(cols):
                    cell = tbl.cell(ri, ci)
                    text = cells[ci] if ci < len(cells) else ""
                    cell.margin_left = cell.margin_right = emu(12)
                    cell.margin_top = cell.margin_bottom = emu(9 if kind == "head" else 15)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    tf = cell.text_frame
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    run = p.add_run()
                    run.text = text
                    if kind == "head":
                        self.style_run(run, "label", self.sz("table_head"), "muted", bold=False)
                    elif ci == 0 and numbered:
                        self.style_run(run, "number", self.sz("table_number"), "accent")
                    elif ci == (1 if numbered else 0):
                        self.style_run(run, "heading", self.sz("table_item"))
                    else:
                        self.style_run(run, "body", self.sz("table"), bold=False)
                    if kind == "head" and head_fill:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = self.rgb("surface")
                    else:
                        cell.fill.background()
                    self._cell_border(cell, "B", "rule" if kind == "row" else
                                      ("rule" if head_fill else "rule_mid"), 0.75 if kind == "row" or head_fill else 1.5)
        self.footer(slide, number)
        self.notes(slide, spec)

    def _cell_border(self, cell, side: str, color: str, weight: float) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        order = ["lnL", "lnR", "lnT", "lnB"]
        for name in order:
            existing = tcpr.find(qn(f"a:{name}"))
            if existing is not None:
                tcpr.remove(existing)
        lines = {}
        for name in order:
            ln = etree.Element(qn(f"a:{name}"))
            if name == f"ln{side}":
                ln.set("w", str(int(emu(weight))))
                fill = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(fill, qn("a:srgbClr")).set("val", self.t.hex(color))
            else:
                ln.set("w", "0")
                etree.SubElement(ln, qn("a:noFill"))
            lines[name] = ln
        for name in reversed(order):
            tcpr.insert(0, lines[name])

    # ------------------------------------------------------------ 大数字页、引文页、图片页

    def number(self, spec: Dict[str, Any], number: int) -> None:
        slide = self.new_slide()
        label = str(spec.get("label") or spec.get("title") or "")
        top = self.label_row(slide, label) if label else self.mt
        bottom = H - self.mb - self.sz("footer") * 1.3 - 24 - 12
        side_w = 255.0 if (spec.get("text") or spec.get("chart")) else 0
        main_w = W - self.ml - self.mr - (side_w + 48 if side_w else 0)
        value, unit = str(spec.get("value", "")), str(spec.get("unit", ""))
        big = self.fit(value, self.sz("big_number"), main_w * 0.72, "display", lines=1,
                       minimum=self.sz("big_number") * 0.4)
        delta = str(spec.get("delta", ""))
        delta_size = self.sz("big_delta")
        y_delta = bottom - delta_size * 1.1
        if delta:
            runs = []
            if spec.get("delta_label"):
                runs.append(r(str(spec["delta_label"]) + "  ", role="label", size=self.sz("big_label") * 1.2,
                              color="muted"))
            runs.append(r(delta, role="display", size=delta_size, color="accent", bold=False))
            self.box(slide, self.ml, y_delta, main_w, delta_size * 1.2, [para(*runs, leading=1.0)], anchor="bottom")
        y_number = (y_delta - 15 if delta else bottom) - big * 0.86
        runs = [r(value, role="display", size=big, bold=False)]
        if unit:
            runs.append(r(" " + unit, role="title", size=self.sz("big_unit")))
        self.box(slide, self.ml, y_number, main_w, big * 0.92, [para(*runs, leading=0.86)], anchor="bottom")
        if side_w:
            sx = W - self.mr - side_w
            y = max(top + 30, bottom - 220)
            if spec.get("text"):
                text = str(spec["text"])
                size = self.sz("big_text")
                h = self.height(text, size, 1.8, side_w, "body")
                self.box(slide, sx, y, side_w, h + 4, [para(r(text, role="body", size=size, color="muted"),
                                                             leading=1.8)])
                y += h + 21
            chart = spec.get("chart")
            if chart and chart.get("data"):
                self._stacked(slide, chart, sx, y, side_w)
        self.footer(slide, number)
        self.notes(slide, spec)

    def _stacked(self, slide, chart: Dict[str, Any], x: float, y: float, width: float) -> None:
        data = [(str(row[0]), float(row[1])) for row in chart["data"]]
        total = sum(v for _, v in data) or 1
        colors = list(self.t.chart_colors())
        highlight = chart.get("highlight")
        if chart.get("title"):
            self.box(slide, x, y, width, 16, [para(r(str(chart["title"]), role="label", size=self.sz("chart_note"),
                                                   color="muted"))])
            y += 21
        cx = x
        palette = []
        for i, (name, value) in enumerate(data):
            color = colors[i % len(colors)]
            if highlight is not None and (highlight == i or highlight == name):
                color = self.t.color("accent")
            palette.append(color)
            seg = (width - 1.5 * (len(data) - 1)) * value / total
            if seg > 0:
                self.rect(slide, cx, y, seg, 10.5, color)
            cx += seg + 1.5
        y += 10.5 + 12
        col_w = width / 2
        for i, (name, value) in enumerate(data):
            lx, ly = x + (i % 2) * col_w, y + (i // 2) * 18
            self.rect(slide, lx, ly + 3, 7.5, 7.5, palette[i])
            text = f"{name} {value:g}"
            on = highlight is not None and (highlight == i or highlight == name)
            self.box(slide, lx + 13, ly, col_w - 13, 16,
                     [para(r(text, role="label", size=self.sz("chart_note"), color="accent" if on else "ink",
                             bold=False))])

    def quote(self, spec: Dict[str, Any], number: int) -> None:
        slide = self.new_slide()
        bottom = H - self.mb - self.sz("footer") * 1.3 - 24
        mark = self.sz("quote_mark")
        self.box(slide, self.ml, self.mt + 54, 110, mark * 0.9,
                 [para(r("“", role="display", size=mark, color="accent", bold=False), leading=0.8)])
        x = self.ml + 112
        width = W - self.mr - x
        text = str(spec.get("quote") or spec.get("content") or spec.get("title") or "")
        size = self.fit(text, self.sz("quote"), width, "quote", lines=4, minimum=self.sz("quote") * 0.5)
        text_h = self.height(text, size, 1.45, width, "quote")
        author = str(spec.get("author", ""))
        meta = [str(spec[k]) for k in ("role", "source") if spec.get(k)]
        attribution_h = (2 + 10.5 + (self.sz("quote_name") * 1.4 if author else 0)
                         + (self.sz("quote_role") * 1.5 + 6 if meta else 0)) if (author or meta) else 0
        block = text_h + (30 + attribution_h if attribution_h else 0)
        y = self.mt + max(0, (bottom - self.mt - block) / 2)
        self.box(slide, x, y, width, text_h + 6, [para(r(text, role="quote", size=size, bold=False), leading=1.45)])
        if attribution_h:
            y += text_h + 30
            self.rect(slide, x, y, 48, 1.5, "accent")
            y += 12
            if author:
                self.box(slide, x, y, width, self.sz("quote_name") * 1.4,
                         [para(r(author, role="title", size=self.sz("quote_name")))])
                y += self.sz("quote_name") * 1.4 + 6
            if meta:
                self.box(slide, x, y, width, self.sz("quote_role") * 1.5,
                         [para(r("　　".join(meta), role="label", size=self.sz("quote_role"), color="muted",
                                 bold=False))])
        self.footer(slide, number)
        self.notes(slide, spec)

    def image(self, spec: Dict[str, Any], number: int) -> None:
        slide = self.new_slide()
        top = self.heading(slide, spec) + 24 if spec.get("title") else self.mt
        bottom = H - self.mb - self.sz("footer") * 1.3 - 24
        caption = str(spec.get("caption", ""))
        cap_h = self.sz("small") * 1.6 if caption else 0
        src = str(spec.get("image", ""))
        path = src if os.path.isabs(src) else os.path.join(self.base_dir, src)
        if not os.path.isfile(path):
            self.warnings.append({"code": "IMAGE_NOT_FOUND", "message": f"图片不存在：{src}"})
            self.box(slide, self.ml, top, W - self.ml - self.mr, 24,
                     [para(r(f"[图片：{caption or src}]", role="body", size=self.sz("body"), color="muted"))])
        else:
            from PIL import Image
            with Image.open(path) as img:
                iw, ih = img.size
            box_w, box_h = W - self.ml - self.mr, bottom - top - cap_h - 6
            scale = min(box_w / iw, box_h / ih)
            w, h = iw * scale, ih * scale
            slide.shapes.add_picture(path, emu(self.ml), emu(top), emu(w), emu(h))
            if caption:
                self.box(slide, self.ml, top + h + 6, box_w, cap_h,
                         [para(r(caption, role="label", size=self.sz("small"), color="muted", bold=False))])
        self.footer(slide, number)
        self.notes(slide, spec)

    def notes(self, slide, spec: Dict[str, Any]) -> None:
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])

    # ------------------------------------------------------------ 入口

    def render(self, slides: List[Dict[str, Any]]) -> None:
        for index, spec in enumerate(slides, start=1):
            layout = str(spec.get("layout") or "content")
            if layout not in LAYOUTS:
                self.warnings.append({"code": "UNKNOWN_LAYOUT",
                                      "message": f"第 {index} 页的版式 {layout} 不存在，已按 content 处理"})
                layout = "content"
            handler = {"title": self.cover}.get(layout) or getattr(self, layout)
            handler(spec, index)


def render_slides(title: Optional[str], slides: List[Dict[str, Any]], output_path: str, theme: Any = None,
                  base_dir: Optional[str] = None, font_policy: str = "local",
                  footer_label: Optional[str] = None, **meta: Any) -> Dict[str, Any]:
    """按主题生成 PPT。meta 可传 brand、classification、publication、issue、season、date、author、kicker 等。"""
    theme = load_theme(theme)
    info = {k: v for k, v in meta.items() if v is not None}
    if title:
        info["title"] = title
    info.setdefault("title", "")
    renderer = SlideRenderer(theme, info, base_dir or os.getcwd(), font_policy, footer_label)
    renderer.render(slides)
    renderer.prs.core_properties.title = str(info.get("title", ""))
    renderer.prs.save(output_path)
    result: Dict[str, Any] = {"output_path": output_path, "file_size": os.path.getsize(output_path),
                              "slides_count": len(renderer.prs.slides), "theme": theme.name}
    if renderer.warnings:
        result["warnings"] = renderer.warnings
    return result
