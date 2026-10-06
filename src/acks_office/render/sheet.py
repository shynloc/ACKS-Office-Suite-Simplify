"""主题驱动的 Excel。

每张工作表：标题行（可选）、表头（fill 浅底加粗 / label 小号标签）、数据行之间的细线、合计行
（rules 上下各一条细线 / bar 上方一条粗线）、数字格式（千位分隔，负数加括号并用负数色）、
隐藏网格线、冻结表头、工作表标签用强调色；可选图表。
"""

import json
import math
import os
import re
import tempfile
import zipfile
from typing import Any, Dict, List, Optional, Sequence

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.drawing.text import Font as TextFont
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..themes import Theme, load_theme
from ..themes.model import one_line
from .common import display_width, is_total_row, numeric_text

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_DASHES = ("-", "—", "–", "/")
# 年份、编号这类列：不加千位分隔，按文字左对齐
_PLAIN = re.compile(r"年份|年度|^年$|编号|序号|代码|工号|邮编|^(id|no\.?|code|year)$", re.I)
_PAPER = {"A4": 9, "A5": 11, "Letter": 1}
# 右对齐单元格的右边距：数字格式每段末尾留一个数字宽的空白。不用缩进——LibreOffice 遇到相邻单元格只差数字格式、
# 边框或条件格式时会丢掉右对齐的缩进，同一列时有时无；Excel 和 LibreOffice 都认数字格式里的空白
_PAD = "_0"


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _clean(value: Any) -> Any:
    """写进单元格前整理：NaN / 无穷大留空，列表和字典转成 JSON 文字，去掉 Excel 不允许的控制字符。"""
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (list, dict, tuple, set)):
        value = json.dumps(list(value) if isinstance(value, (tuple, set)) else value, ensure_ascii=False)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _decimals(value: float) -> int:
    return len(f"{value:.10f}".rstrip("0").split(".")[1])


def number_format(values: Sequence[Any]) -> Optional[str]:
    """一列数值的数字格式：千位分隔，小数按数据保留 0–2 位，负数显示为括号。

    正数、零和文字后面留出一个右括号的宽度（会计格式的 _)），和负数的数字对齐。
    """
    numbers = [float(v) for v in values if _is_number(v) and math.isfinite(v)]
    if not numbers:
        return None
    decimals = min(max(_decimals(v) for v in numbers), 2)
    base = "#,##0" + ("." + "0" * decimals if decimals else "")
    return f"{base}_);({base});{base}_);@_)"


def _sections(fmt: str) -> List[str]:
    """按分号拆开数字格式的各段（引号里和反斜杠转义的分号不算）。"""
    parts, current, quoted, escaped = [], "", False, False
    for ch in fmt:
        if escaped:
            escaped = False
        elif ch == "\\":
            escaped = True
        elif ch == '"':
            quoted = not quoted
        elif ch == ";" and not quoted:
            parts.append(current)
            current = ""
            continue
        current += ch
    return parts + [current]


def padded(fmt: str) -> str:
    """数字格式每段末尾加上右边距（空的段落保持为空，仍然不显示）。"""
    return ";".join(section + _PAD if section else section for section in _sections(fmt))


def _shown(value: Any, fmt: Optional[str]) -> str:
    """单元格大致会显示成什么样，用来估算列宽。"""
    if value is None or (isinstance(value, str) and value.startswith("=")):
        return ""
    if not (_is_number(value) and fmt):
        return str(value)
    section = _sections(fmt)[0]
    decimals = len(re.match(r"[0#]*", section.partition(".")[2]).group(0))
    percent = "%" in section
    text = f"{abs(value) * (100 if percent else 1):,.{decimals}f}" + ("%" if percent else "")
    return f"({text})"  # 给括号或正负号留出位置


def sheet_name(name: str, used: List[str]) -> str:
    """合法且不重名的工作表名：去掉 []:*?/\\，最长 31 个字符。"""
    clean = _BAD_SHEET_CHARS.sub("", one_line(str(name or ""))).strip(" '") or "Sheet"
    clean = clean[:31]
    candidate, n = clean, 2
    while candidate.lower() in (u.lower() for u in used):
        suffix = f" ({n})"
        candidate, n = clean[:31 - len(suffix)] + suffix, n + 1
    used.append(candidate)
    return candidate


class SheetRenderer:
    def __init__(self, theme: Theme, policy: str = "local"):
        self.t = theme
        self.layout = theme.get("layout.sheet")
        self.warnings: List[Dict[str, str]] = []
        self.cached: Dict[str, Dict[str, float]] = {}  # 工作表名 → {单元格: 合计公式的计算结果}
        self.fonts = {role: theme.font(role, policy) for role in ("body", "label", "title")}
        for choice in self.fonts.values():
            for note in choice.substitutions:
                if not any(w.get("family") == note["family"] for w in self.warnings):
                    self.warnings.append(note)

    # ------------------------------------------------------------ 样式

    def sz(self, key: str) -> float:
        return float(self.t.get(f"type.sheet.{key}"))

    def font(self, role: str, size: float, color: str = "ink", bold: Optional[bool] = None) -> Font:
        # 一个单元格只能设一种字体：用中文字体（中文字体也带西文字形），中西文都能正常显示
        choice = self.fonts[role]
        return Font(name=choice.cn, size=size, color=self.t.hex(color), bold=choice.bold if bold is None else bold)

    def side(self, color: str, style: str = "thin") -> Side:
        return Side(style=style, color=self.t.hex(color))

    @staticmethod
    def align(right: bool) -> Alignment:
        # 左对齐用缩进留出左边距；右对齐的右边距写在数字格式里（见 _PAD）
        return Alignment(horizontal="right" if right else "left", vertical="center", indent=0 if right else 1)

    def height(self, factor: float) -> float:
        return round(self.sz("body") * factor, 1)

    # ------------------------------------------------------------ 工作表

    def render(self, wb, specs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        used: List[str] = []
        summary = []
        for i, spec in enumerate(specs):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = sheet_name(spec.get("name") or spec.get("title") or f"Sheet{i + 1}", used)
            summary.append(self.sheet(ws, spec))
        return summary

    def sheet(self, ws, spec: Dict[str, Any]) -> Dict[str, Any]:
        header = ["" if h is None else str(_clean(h)) for h in spec.get("header") or []]
        rows = [[_clean(v) for v in row] for row in spec.get("rows") or []]
        width = max([len(header)] + [len(row) for row in rows])
        header += [""] * (width - len(header))
        rows = [row + [None] * (width - len(row)) for row in rows]
        columns = self._columns(header, rows, spec.get("formats") or {})

        r = 1
        if spec.get("title"):
            cell = ws.cell(row=1, column=1, value=one_line(str(_clean(spec["title"]))))
            cell.font = self.font("title", self.sz("title"))
            cell.alignment = Alignment(vertical="center", indent=1)
            ws.row_dimensions[1].height = self.height(3.0)
            r = 2
        head_row = None
        if any(header):
            head_row = r
            self._header(ws, r, header, columns)
            r += 1

        first = r
        total = spec.get("total")
        detected = total is not False and bool(rows) and is_total_row(["" if rows[-1][0] is None else str(rows[-1][0])])
        appended = total is True and bool(rows) and not detected
        count = len(rows) + (1 if appended else 0)
        total_index = count - 1 if (detected or appended) else None
        for i, values in enumerate(rows):
            self._row(ws, r, values, columns, i == total_index, self._border(i, count, total_index))
            r += 1
        if appended:
            values: List[Any] = [self.t.label("total")]
            for c in range(1, width):
                letter = get_column_letter(c + 1)
                summed = [row[c] for row in rows]
                if not columns[c]["numbers"] or columns[c]["plain"]:
                    values.append(None)
                    continue
                values.append(f"=SUM({letter}{first}:{letter}{r - 1})")
                if all(v is None or _is_number(v) for v in summed):  # 都是数字才能替公式算出结果
                    total_value = round(math.fsum(v for v in summed if _is_number(v)), 9)
                    self.cached.setdefault(ws.title, {})[f"{letter}{r}"] = total_value
            self._row(ws, r, values, columns, True, self._border(count - 1, count, total_index))
            r += 1
        last = r - 1

        self._negatives(ws, first, last, columns)
        for item in spec.get("highlight") or []:
            self._highlight(ws, item, first, count, width)
        if spec.get("note"):  # 紧接表格，用行高留出间距：读取数据时不会多出一行空行
            cell = ws.cell(row=last + 1, column=1, value=str(_clean(spec["note"])))
            cell.font = self.font("label", self.sz("note"), "muted", bold=False)
            cell.alignment = Alignment(vertical="bottom", indent=1)
            ws.row_dimensions[last + 1].height = self.height(3.0)
        self._widths(ws, header, rows, columns)

        ws.sheet_view.showGridLines = bool(self.layout["gridlines"])
        if self.layout["freeze_header"] and head_row:
            ws.freeze_panes = ws.cell(row=head_row + 1, column=1).coordinate
        ws.sheet_properties.tabColor = self.t.hex(self.layout["tab_color"])
        self._print_setup(ws, width, head_row)
        if spec.get("chart") is not None and spec.get("chart") is not False:
            chart = spec["chart"] if isinstance(spec["chart"], dict) else {}
            below = last + (3 if spec.get("note") else 2)  # 图表放在表格下方：打印时不用为了并排而缩小
            self._chart(ws, chart, header, columns, head_row, first, last - (1 if total_index is not None else 0), below)
        return {"name": ws.title, "rows": count, "columns": width}

    def _columns(self, header: List[str], rows: List[List[Any]], formats: Dict[Any, str]) -> List[Dict[str, Any]]:
        """每列的数字格式、是否按数字右对齐、是否是年份 / 编号这类不加千位分隔的列。"""
        formats = {str(k): v for k, v in formats.items()}
        result = []
        for c, name in enumerate(header):
            values = [row[c] for row in rows]
            filled = [v for v in values if v is not None and str(v).strip() and str(v).strip() not in _DASHES]
            plain = bool(name and _PLAIN.search(name.strip()))
            fmt = formats.get(str(c)) or (formats.get(name) if name else None) or ("0" if plain else number_format(values))
            numeric = bool(filled) and not plain and all(
                _is_number(v) or (isinstance(v, str) and (numeric_text(v) or v.startswith("="))) for v in filled)
            column = {"format": fmt, "raw": fmt, "numeric": numeric, "plain": plain, "head": None,
                      "numbers": any(_is_number(v) for v in values)}
            if numeric:  # 表头和数字一样靠右，右边距相同；会计格式的列再让出右括号的宽度
                column["head"] = padded("@_)" if fmt and _sections(fmt)[0].endswith("_)") else "@")
                column["format"] = padded(fmt or "General;-General;General;@")
            result.append(column)
        return result

    def _header(self, ws, r: int, header: List[str], columns: List[Dict[str, Any]]) -> None:
        label = self.layout["head"] == "label"
        for c, name in enumerate(header):
            cell = ws.cell(row=r, column=c + 1, value=name or None)
            if label:
                cell.font = self.font("label", self.sz("head"), "muted", bold=False)
            else:
                cell.font = self.font("label", self.sz("head"), bold=True)
                cell.fill = PatternFill("solid", fgColor=self.t.hex("surface"))
            cell.alignment = self.align(columns[c]["numeric"])
            if columns[c]["head"]:
                cell.number_format = columns[c]["head"]
            cell.border = Border(bottom=self.side(self.layout["head_rule"], "medium"))
        ws.row_dimensions[r].height = self.height(2.6)

    def _row(self, ws, r: int, values: List[Any], columns: List[Dict[str, Any]], total: bool, border: Border) -> None:
        for c, value in enumerate(values):
            cell = ws.cell(row=r, column=c + 1, value=value)
            cell.font = self.font("body", self.sz("body"), bold=total)
            fmt = columns[c]["format"]
            if fmt and (columns[c]["numeric"] or _is_number(value) or (isinstance(value, str) and value.startswith("="))):
                cell.number_format = fmt
            cell.alignment = self.align(columns[c]["numeric"])
            cell.border = border
        ws.row_dimensions[r].height = self.height(2.4)

    def _border(self, i: int, count: int, total_index: Optional[int]) -> Border:
        """第 i 行数据（共 count 行，含合计行）的边框。"""
        sheet = self.layout
        bar = sheet["total"] == "bar"
        if i == total_index:
            if bar:  # 上方一条粗线
                return Border(top=self.side(sheet["total_rule"], "medium"))
            return Border(top=self.side(sheet["total_rule"]), bottom=self.side(sheet["total_rule"]))
        if total_index is not None and i == total_index - 1:
            return Border()  # 合计行上方的线已经把它和数据隔开
        if i == count - 1:  # 没有合计行时，最后一行收尾
            return Border(bottom=self.side(sheet["row_rule"] if bar else sheet["total_rule"]))
        return Border(bottom=self.side(sheet["row_rule"]))

    def _negatives(self, ws, first: int, last: int, columns: List[Dict[str, Any]]) -> None:
        """负数用主题的负数色（条件格式，数字格式本身只负责括号）。"""
        if last < first:
            return
        color = Font(color=self.t.hex(self.layout["negative"]))
        for c, column in enumerate(columns):
            if (column["numeric"] or column["numbers"]) and not column["plain"]:  # 数字列和含数字、公式的列
                letter = get_column_letter(c + 1)
                ws.conditional_formatting.add(f"{letter}{first}:{letter}{last}",
                                              CellIsRule(operator="lessThan", formula=["0"], font=color))

    def _highlight(self, ws, item: Any, first: int, count: int, width: int) -> None:
        """重点单元格加一圈粗框；位置为 [数据行, 列]，都从 0 数起。"""
        try:
            i, c = int(item[0]), int(item[1])
        except (TypeError, ValueError, IndexError):
            i = c = -1
        if not (0 <= i < count and 0 <= c < width):
            self.warnings.append({"code": "BAD_HIGHLIGHT", "message": f"{ws.title}：重点单元格位置无效 {item!r}，"
                                                                      f"应为 [数据行, 列]，从 0 数起"})
            return
        edge = self.side(self.layout["highlight"], "medium")
        ws.cell(row=first + i, column=c + 1).border = Border(left=edge, right=edge, top=edge, bottom=edge)

    def _widths(self, ws, header: List[str], rows: List[List[Any]], columns: List[Dict[str, Any]]) -> None:
        head_size = self.sz("head") * (1.08 if self.layout["head"] == "fill" else 1.0)  # 加粗略宽
        for c, column in enumerate(columns):
            sized = [(header[c], head_size)] + [(_shown(row[c], column["format"]), self.sz("body")) for row in rows]
            widest = max([display_width(text) * size / 11 * 1.15 for text, size in sized if text] + [0])
            units = min(max(widest + 3.5, 10 if column["numeric"] else 8), 60)
            ws.column_dimensions[get_column_letter(c + 1)].width = round(units, 1)

    def _print_setup(self, ws, width: int, head_row: Optional[int]) -> None:
        """打印：按主题的纸张大小，宽度缩到一页，每页重复表头。"""
        ws.page_setup.orientation = "landscape" if width > 6 else "portrait"
        ws.page_setup.paperSize = _PAPER.get(self.t.get("layout.doc.page.size", "A4"), 9)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if head_row:
            ws.print_title_rows = f"{head_row}:{head_row}"

    # ------------------------------------------------------------ 图表

    def _text(self, size: float, color: str = "muted", bold: bool = False) -> CharacterProperties:
        choice = self.fonts["label"]
        return CharacterProperties(latin=TextFont(typeface=choice.en), ea=TextFont(typeface=choice.cn),
                                   sz=int(size * 100), b=bold, solidFill=self.t.hex(color))

    def _rich(self, size: float, color: str = "muted") -> RichText:
        props = self._text(size, color)
        return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props)])

    def _chart(self, ws, spec: Dict[str, Any], header: List[str], columns: List[Dict[str, Any]],
               head_row: Optional[int], first: int, data_last: int, below: int) -> None:
        """柱状图 / 条形图 / 折线图：分类取第一列，数值列默认取第一个数字列，颜色用主题的图表色。"""
        values = spec.get("values")
        if values is None:
            values = [c for c, col in enumerate(columns) if c > 0 and col["numbers"] and not col["plain"]][:1]
        values = [int(v) for v in values if isinstance(v, int) and 0 < v < len(columns)]
        if not values or data_last < first:
            self.warnings.append({"code": "CHART_SKIPPED", "message": f"{ws.title}：没有可画图的数值列，未生成图表"})
            return
        kind = spec.get("type", "column")
        chart = LineChart() if kind == "line" else BarChart()
        if kind != "line":
            chart.type = "bar" if kind == "bar" else "col"
            chart.gapWidth = 60
        colors = self.t.chart_colors()
        for i, c in enumerate(values):
            top = head_row if head_row == first - 1 else first
            chart.add_data(Reference(ws, min_col=c + 1, min_row=top, max_row=data_last), titles_from_data=top != first)
            series = chart.series[-1]
            color = colors[i % len(colors)].lstrip("#")
            if kind == "line":
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 28575  # 2.25 磅
                series.smooth = False
            else:
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.noFill = True
        chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=data_last))
        title = spec.get("title") or (header[values[0]] if len(values) == 1 else None)
        if title:
            props = self._text(self.sz("body"), "ink", True)
            run = RegularTextRun(rPr=props, t=str(title))
            chart.title = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), r=[run])])),
                                overlay=False)
        if len(values) == 1:
            chart.legend = None
        else:
            chart.legend.position = "b"
            chart.legend.txPr = self._rich(self.sz("note"))
        for axis in (chart.x_axis, chart.y_axis):
            axis.delete = False
            axis.majorGridlines = None
            axis.txPr = self._rich(self.sz("note"))
        chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=self.t.hex("rule_mid")))
        chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        if columns[values[0]]["raw"]:  # 坐标轴用不带右边距的格式
            chart.y_axis.number_format = columns[values[0]]["raw"]
        chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
        chart.roundedCorners = False
        chart.width, chart.height = 16, 8.5
        ws.add_chart(chart, spec.get("anchor") or f"A{below}")


def render_sheets(title: Optional[str], output_path: str, data: Optional[List[List[Any]]] = None,
                  sheets: Optional[List[Dict[str, Any]]] = None, theme: Any = None,
                  font_policy: str = "local", create_chart: bool = False, sheet_name: Optional[str] = None,
                  **meta: Any) -> Dict[str, Any]:
    """按主题生成 Excel。

    data：二维数组，第一行为表头（单张工作表，标题行为 title，工作表名为 sheet_name 或 title）；
    sheets：多张工作表，每项 {name, title, header, rows, formats, total, highlight, note, chart}；
    create_chart：没写 chart 的工作表也画一张默认图表。
    """
    theme = load_theme(theme)
    if sheets is None:
        rows = [list(r) for r in data or []]
        sheets = [{"name": sheet_name or title, "title": title, "header": rows[0] if rows else [], "rows": rows[1:]}]
    elif not (isinstance(sheets, list) and sheets and all(isinstance(spec, dict) for spec in sheets)):
        raise ValueError("sheets 需要是非空数组，每项是一张工作表 {name, header, rows, …}")
    sheets = [dict(spec) for spec in sheets]
    if create_chart:
        for spec in sheets:
            spec.setdefault("chart", {})
    renderer = SheetRenderer(theme, font_policy)
    wb = openpyxl.Workbook()
    summary = renderer.render(wb, sheets)
    wb.properties.title = one_line(str(title or ""))
    if meta.get("author"):
        wb.properties.creator = str(meta["author"])
    wb.save(output_path)
    if renderer.cached:
        _store_cached(output_path, [ws.title for ws in wb.worksheets], renderer.cached)
    result: Dict[str, Any] = {"output_path": output_path, "file_size": os.path.getsize(output_path),
                              "sheets": summary, "rows": summary[0]["rows"], "columns": summary[0]["columns"],
                              "theme": theme.name}
    if renderer.warnings:
        result["warnings"] = renderer.warnings
    return result


def _store_cached(path: str, titles: List[str], cached: Dict[str, Dict[str, float]]) -> None:
    """openpyxl 不写公式的计算结果：把合计行 SUM 的结果补进文件，pandas 等直接读文件时也能读到合计。

    Excel、LibreOffice 打开时照常重算（openpyxl 写了 fullCalcOnLoad）。
    """
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(path)))
    os.close(fd)
    try:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as target:
            for info, data in entries:
                match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
                values = cached.get(titles[int(match.group(1)) - 1]) if match else None
                if values:
                    xml = data.decode("utf-8")
                    for ref, value in values.items():
                        xml = re.sub(rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)(?:<v>\s*</v>|<v\s*/>)?',
                                     lambda m, v=value: f"{m.group(1)}<v>{v!r}</v>", xml, count=1)
                    data = xml.encode("utf-8")
                target.writestr(info, data)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
