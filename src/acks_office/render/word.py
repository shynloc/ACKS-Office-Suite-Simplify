"""主题驱动的 Word 渲染。

版式由主题令牌决定：封面（title-block 标题块 / standard 标准封面 / issue 期刊封面）、
章节（inline 编号标题 / opener 章节页）、栏数、首字下沉、提示块（tint / rule）、引文（bar / pull）、
表格表头（fill / label）。正文用 Word 命名样式（标题 1–3、正文、说明等），便于之后编辑。
"""

import os
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.text.run import Run

from .. import markdown_blocks as mdb
from ..themes import Theme, load_theme
from . import ooxml as ox
from .common import attach_captions, display_width, document_meta, is_total_row, numeric_columns

PAGE_SIZES = {"A4": (21.0, 29.7), "A5": (14.8, 21.0), "Letter": (21.59, 27.94)}
ROLES = ("body", "title", "heading", "label", "number", "display", "quote", "code")
BULLETS = ("•", "◦", "▪")


class WordRenderer:
    def __init__(self, theme: Theme, meta: Dict[str, Any], base_dir: str, policy: str = "local",
                 footer_label: Optional[str] = None):
        self.t = theme
        self.meta = meta
        self.base_dir = base_dir
        self.footer_label = footer_label
        self.warnings: List[Dict[str, str]] = []
        self.fonts = {role: theme.font(role, policy) for role in ROLES}
        for choice in self.fonts.values():
            for note in choice.substitutions:
                if not any(w.get("family") == note["family"] for w in self.warnings):
                    self.warnings.append(note)
        self.doc = Document()
        self.values = theme.chrome_values(meta)
        self.layout = theme.get("layout.doc")
        self.columns = int(self.layout["columns"])
        self.current_cols = 1
        self.table_no = self.figure_no = self.chapter_no = 0
        self._ids = ox.iter_ids(1)
        self._numbering: Dict[str, int] = {}
        self._fresh_page = True
        self._drop_cap_next = False
        self._setup_page(self.doc.sections[0])
        self._setup_styles()

    # ------------------------------------------------------------ 尺寸与字体

    def sz(self, key: str) -> float:
        return float(self.t.get(f"type.doc.{key}"))

    def lead(self, key: str) -> float:
        return float(self.t.get(f"leading.{key}"))

    def space(self, key: str) -> float:
        return float(self.t.get(f"space.{key}"))

    @property
    def text_width(self) -> float:
        """版心宽度（磅）。"""
        section = self.doc.sections[-1]
        return (section.page_width - section.left_margin - section.right_margin) / 12700

    @property
    def column_width(self) -> float:
        gap = float(self.layout["column_gap_cm"]) * 28.3465
        return (self.text_width - gap * (self.current_cols - 1)) / self.current_cols

    def rgb(self, color: str) -> RGBColor:
        return RGBColor.from_string(self.t.hex(color))

    def run(self, run: Run, role: str = "body", size: Optional[float] = None, color: str = "ink",
            bold: Optional[bool] = None, italic: bool = False, track: float = 0) -> Run:
        """给一段文字设字体角色、字号、颜色；track 为字距（em）。"""
        choice = self.fonts[role]
        ox.run_fonts(run._r, choice.en, choice.cn)
        if size:
            run.font.size = Pt(size)
        run.font.color.rgb = self.rgb(color)
        run.font.bold = choice.bold if bold is None else bold
        if italic:
            run.font.italic = True
        if track and size:
            ox.tracking(run._r.get_or_add_rPr(), track * size)
        return run

    def text(self, paragraph, text: str, **style) -> Run:
        return self.run(paragraph.add_run(text), **style)

    def fmt(self, paragraph, size: float, leading: float, before: float = 0, after: float = 0,
            exact: bool = False, align=None, left: float = 0, right: float = 0, first: float = 0):
        pf = paragraph.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY if exact else WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt(size * leading)
        if align is not None:
            pf.alignment = align
        if left:
            pf.left_indent = Pt(left)
        if right:
            pf.right_indent = Pt(right)
        if first:
            pf.first_line_indent = Pt(first)
        return paragraph

    # ------------------------------------------------------------ 页面与样式

    def _setup_page(self, section) -> None:
        width, height = PAGE_SIZES[self.layout["page"]["size"]]
        margin = self.layout["page"]["margin_cm"]
        section.page_width, section.page_height = Cm(width), Cm(height)
        section.top_margin, section.bottom_margin = Cm(margin["top"]), Cm(margin["bottom"])
        section.left_margin, section.right_margin = Cm(margin["left"]), Cm(margin["right"])
        section.header_distance = Cm(self.layout["page"]["header_cm"])
        section.footer_distance = Cm(self.layout["page"]["footer_cm"])

    def _style(self, name: str, role: str, size: float, color: str = "ink", leading: float = 1.3,
               before: float = 0, after: float = 0, exact: bool = False, bold: Optional[bool] = None,
               italic: bool = False, align=None, keep_next: bool = False, base: str = "Normal"):
        styles = self.doc.styles
        try:
            style = styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles[base]
        style.hidden = False
        choice = self.fonts[role]
        rpr = style.element.get_or_add_rPr()
        ox.run_fonts(rpr, choice.en, choice.cn)
        style.font.size = Pt(size)
        style.font.color.rgb = self.rgb(color)
        style.font.bold = choice.bold if bold is None else bold
        style.font.italic = italic
        pf = style.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY if exact else WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt(size * leading)
        if align is not None:
            pf.alignment = align
        pf.keep_with_next = keep_next
        ppr = style.element.get_or_add_pPr()
        for tag in ("w:pBdr", "w:shd", "w:tabs"):  # 去掉模板自带的下划线、底纹、制表位
            node = ppr.find(qn(tag))
            if node is not None:
                ppr.remove(node)
        return style

    def _setup_styles(self) -> None:
        justify = WD_ALIGN_PARAGRAPH.JUSTIFY if self.layout["justify"] else WD_ALIGN_PARAGRAPH.LEFT
        body, para = self.sz("body"), self.space("paragraph")
        self._style("Normal", "body", body, leading=self.lead("body"), after=para, align=justify)
        heading = self.lead("heading")
        for level, key in ((1, "h1"), (2, "h2"), (3, "h3")):
            self._style(f"Heading {level}", "heading", self.sz(key), leading=heading, exact=True,
                        before=self.space(f"{key}_before"), after=self.space(f"{key}_after"),
                        align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        self._style("Title", "title", self.sz("title"), leading=self.lead("title"), exact=True,
                    align=WD_ALIGN_PARAGRAPH.LEFT)
        self._style("Subtitle", "body", self.sz("subtitle"), color="muted", leading=1.6,
                    align=WD_ALIGN_PARAGRAPH.LEFT)
        self._style("Caption", "label", self.sz("caption"), color="muted", leading=1.4, before=4, after=6,
                    align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, keep_next=True)
        self._style("Quote", "quote", self.sz("quote"), color="muted", leading=self.lead("quote"), after=para)
        self._style("Lede", "body", self.sz("lede"), color="muted", leading=self.lead("lede"), after=0,
                    align=WD_ALIGN_PARAGRAPH.LEFT)
        self._style("List Item", "body", body, leading=self.lead("body"), after=3, align=justify)
        table_lead = self.lead("table")
        self._style("Table Text", "body", self.sz("table"), leading=table_lead, align=WD_ALIGN_PARAGRAPH.LEFT)
        self._style("Table Head", "label", self.sz("table_head"), color="muted", leading=table_lead,
                    align=WD_ALIGN_PARAGRAPH.LEFT,
                    bold=self.layout["table"]["head"] == "fill" and self.fonts["label"].bold)
        self._style("Callout Text", "body", self.sz("callout"), leading=self.lead("callout"), after=0,
                    align=justify)
        self._style("Code", "code", self.sz("code"), leading=1.45, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        for name in ("Header", "Footer"):
            self._style(name, "label", self.sz("header"), color="muted", leading=1.3, bold=False,
                        align=WD_ALIGN_PARAGRAPH.LEFT)
        if self.layout["heading_numbers"] and self.layout["chapter"] == "inline":
            self._heading_numbering()
        self.doc.core_properties.title = self.meta.get("title", "")
        if self.meta.get("author"):
            self.doc.core_properties.author = str(self.meta["author"])

    # ------------------------------------------------------------ 编号

    def _abstract_num(self, levels: List[Tuple[str, str, Optional[str], float, float]]) -> int:
        """levels: (编号格式, 编号文字, 关联样式, 左缩进, 悬挂缩进)。返回 abstractNumId。"""
        numbering = self.doc.part.numbering_part.element
        ids = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(ids, default=0) + 1
        abstract = ox.el("w:abstractNum", abstractNumId=abstract_id)
        abstract.append(ox.el("w:multiLevelType", val="multilevel"))
        number_font = self.fonts["number"]
        for i, (fmt, text, style, left, hanging) in enumerate(levels):
            lvl = ox.el("w:lvl", ilvl=i)
            lvl.append(ox.el("w:start", val=1))
            lvl.append(ox.el("w:numFmt", val=fmt))
            if style:
                lvl.append(ox.el("w:pStyle", val=style))
            lvl.append(ox.el("w:lvlText", val=text))
            lvl.append(ox.el("w:lvlJc", val="left"))
            ppr = ox.el("w:pPr")
            tabs = ox.el("w:tabs")
            tabs.append(ox.el("w:tab", val="num", pos=ox.twips(left)))
            ppr.append(tabs)
            ppr.append(ox.el("w:ind", left=ox.twips(left), hanging=ox.twips(hanging)))
            lvl.append(ppr)
            rpr = ox.el("w:rPr")
            fonts = ox.el("w:rFonts", ascii=number_font.en, hAnsi=number_font.en, eastAsia=number_font.cn)
            rpr.append(fonts)
            rpr.append(ox.el("w:color", val=self.t.hex("accent")))
            lvl.append(rpr)
            abstract.append(lvl)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        return abstract_id

    def _num(self, abstract_id: int, restart: Optional[Tuple[int, int]] = None) -> int:
        numbering = self.doc.part.numbering_part.element
        num = numbering.add_num(abstract_id)
        if restart:
            num.add_lvlOverride(ilvl=restart[0]).add_startOverride(restart[1])
        return num.numId

    def _heading_numbering(self) -> None:
        widths = (self.sz("h1") * 1.5, self.sz("h2") * 2.3, self.sz("h3") * 2.8)
        abstract = self._abstract_num([("decimal", "%1", "Heading1", widths[0], widths[0]),
                                       ("decimal", "%1.%2", "Heading2", widths[1], widths[1]),
                                       ("decimal", "%1.%2.%3", "Heading3", widths[2], widths[2])])
        num_id = self._num(abstract)
        for level in range(3):
            style = self.doc.styles[f"Heading {level + 1}"]
            ppr = style.element.get_or_add_pPr()
            num_pr = ox.el("w:numPr")
            num_pr.append(ox.el("w:ilvl", val=level))
            num_pr.append(ox.el("w:numId", val=num_id))
            ox.set_child(ppr, num_pr, "pPr")

    def _list_num(self, ordered: bool, level: int, start: int) -> int:
        if not ordered:
            if "bullet" not in self._numbering:
                abstract = self._abstract_num([("bullet", BULLETS[i], None, 18 * (i + 1), 12) for i in range(3)])
                self._numbering["bullet"] = self._num(abstract)
            return self._numbering["bullet"]
        if "ordered" not in self._numbering:
            self._numbering["ordered"] = self._abstract_num(
                [(fmt, f"%{i + 1}.", None, 20 * (i + 1), 16)
                 for i, fmt in enumerate(("decimal", "lowerLetter", "lowerRoman"))])
        return self._num(self._numbering["ordered"], restart=(level, start))

    # ------------------------------------------------------------ 页眉页脚

    def _chrome_paragraph(self, paragraph, path: str, chapter: Optional[str] = None,
                          rule: bool = False, override: Optional[str] = None) -> None:
        values = dict(self.values, chapter=chapter or "")
        left = self.t.template(f"{path}.left", values) if override is None else [("text", override)]
        right = self.t.template(f"{path}.right", values)
        paragraph.style = self.doc.styles["Header" if path.endswith("header") else "Footer"]
        ox.tab_stop(paragraph, self.text_width, "right")
        size = self.sz("header")

        def style_run(run):
            self.run(run, "label", size, "muted", bold=False)

        for i, segments in enumerate((left, right)):
            if i:
                style_run(paragraph.add_run("\t"))
            for kind, value in segments:
                if kind == "text":
                    style_run(paragraph.add_run(value))
                else:
                    ox.add_field(paragraph, "PAGE" if value == "page" else "NUMPAGES", "1", style=style_run)
        if rule:
            ox.paragraph_borders(paragraph, bottom=(0.75, self.t.hex("rule"), 4))

    def _chrome(self, section, chapter: Optional[str] = None, plain_first_page: bool = False) -> None:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        for container in (section.header, section.footer):
            for paragraph in list(container.paragraphs)[1:]:
                ox.remove_paragraph(paragraph)
            for paragraph in container.paragraphs:
                for child in list(paragraph._p):
                    if child.tag != qn("w:pPr"):
                        paragraph._p.remove(child)
        self._chrome_paragraph(section.header.paragraphs[0], "doc.header", chapter, self.layout["header_rule"])
        footer_override = self.footer_label
        self._chrome_paragraph(section.footer.paragraphs[0], "doc.footer", chapter, override=footer_override)
        if plain_first_page:
            section.different_first_page_header_footer = True
            self._chrome_paragraph(section.first_page_footer.paragraphs[0], "doc.footer", chapter,
                                   override=footer_override)

    # ------------------------------------------------------------ 封面

    def _top_row(self, left: Tuple[str, Dict], right: Tuple[str, Dict], after: float = 0):
        p = self.doc.add_paragraph()
        self.fmt(p, 12, 1.2, after=after, align=WD_ALIGN_PARAGRAPH.LEFT)
        ox.tab_stop(p, self.text_width, "right")
        if left[0]:
            self.text(p, left[0], **left[1])
        if right[0]:
            p.add_run("\t")
            self.text(p, right[0], **right[1])
        return p

    def _short_bar(self, width: float, weight: float, color: str, before: float = 0, after: float = 0):
        p = self.doc.add_paragraph()
        self.fmt(p, 2, 1.0, before=before, after=after, exact=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                 right=max(0, self.text_width - width))
        p.add_run("").font.size = Pt(1)
        ox.paragraph_borders(p, bottom=(weight, self.t.hex(color), 0))
        return p

    def _title_lines(self, paragraph, role: str, size: float, color: str = "ink") -> None:
        lines = str(self.meta.get("title", "")).split("\n")
        for i, line in enumerate(lines):
            run = self.text(paragraph, line, role=role, size=size, color=color)
            if i < len(lines) - 1:
                run.add_break()

    def _title_block(self) -> None:
        """neutral：标题放在第一页顶部，不单独占一页。"""
        if self.meta.get("kicker"):
            p = self.doc.add_paragraph()
            self.fmt(p, self.sz("kicker"), 1.4, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
            self.text(p, self.meta["kicker"], role="label", size=self.sz("kicker"), color="accent")
        p = self.doc.add_paragraph(style="Title")
        self._title_lines(p, "title", self.sz("title"))
        if self.meta.get("subtitle"):
            p = self.doc.add_paragraph(self.meta["subtitle"], style="Subtitle")
            p.paragraph_format.space_before = Pt(6)
        info = " · ".join(str(self.meta[k]) for k in self.t.chrome_get("doc.cover_meta", []) if self.meta.get(k))
        p = self.doc.add_paragraph()
        self.fmt(p, self.sz("small"), 1.4, before=8, after=18, align=WD_ALIGN_PARAGRAPH.LEFT)
        if info:
            self.text(p, info, role="label", size=self.sz("small"), color="muted", bold=False)
        ox.paragraph_borders(p, bottom=(0.75, self.t.hex("rule_strong"), 6))
        self._fresh_page = False

    def _cover_standard(self, section) -> None:
        """slate：品牌行、强调短线、类型、标题、副标题；日期 / 编制 / 版本放在页脚，固定在页底。"""
        values = self.values
        self._top_row((values.get("brand", ""), dict(role="heading", size=self.sz("brand"), track=0.02)),
                      (values.get("classification", ""),
                       dict(role="label", size=self.sz("classification"), color="muted", bold=False)))
        self._short_bar(30, 3, "accent", before=self.space("cover_title_top"), after=10)
        if self.meta.get("kicker"):
            p = self.doc.add_paragraph()
            self.fmt(p, self.sz("kicker"), 1.3, after=10, align=WD_ALIGN_PARAGRAPH.LEFT)
            self.text(p, self.meta["kicker"], role="label", size=self.sz("kicker"), color="accent")
        p = self.doc.add_paragraph(style="Title")
        p.paragraph_format.space_after = Pt(10)
        self._title_lines(p, "title", self.sz("title"))
        if self.meta.get("subtitle"):
            self.doc.add_paragraph(self.meta["subtitle"], style="Subtitle")
        self._cover_footer_meta(section)

    def _cover_footer_meta(self, section) -> None:
        rows = [(self.t.label(k), str(self.meta[k])) for k in self.t.chrome_get("doc.cover_meta", [])
                if self.meta.get(k)]
        footer = section.footer
        footer.is_linked_to_previous = False
        if not rows:
            return
        table = footer.add_table(rows=len(rows), cols=2, width=Pt(self.text_width))
        self._plain_table(table, [54, self.text_width - 54])
        for i, (label, value) in enumerate(rows):
            for j, (text, style) in enumerate(((label, dict(role="label", size=self.sz("meta_label"),
                                                             color="muted", bold=False)),
                                               (value, dict(role="body", size=self.sz("meta_value"))))):
                cell = table.cell(i, j)
                ox.cell_margins(cell, 7, 0, 7, 0)
                p = cell.paragraphs[0]
                self.fmt(p, style["size"], 1.3, align=WD_ALIGN_PARAGRAPH.LEFT)
                self.text(p, text, **style)
                sides = {"bottom": (0.75, self.t.hex("rule"))}
                if i == 0:
                    sides["top"] = (0.75, self.t.hex("rule_strong"))
                ox.cell_borders(cell, **sides)
        first = footer.paragraphs[0]
        first._p.getparent().remove(first._p)
        footer._element.append(first._p)  # 页脚须以段落结尾：把原来的空段落移到表格之后
        self.fmt(first, 2, 1.0, exact=True)

    def _cover_issue(self, section, chapters: List[Dict[str, str]]) -> None:
        """folio：刊名与季节、超大期号、标题、导语、朱砂短线；本期目录放在页脚，固定在页底。"""
        values = self.values
        self._top_row((values.get("publication") or values.get("brand", ""),
                       dict(role="label", size=self.sz("brand"), track=self.t.get("tracking.label"))),
                      (str(self.meta.get("season", "")).upper(),
                       dict(role="display", size=self.sz("classification"), color="muted", bold=True,
                            track=self.t.get("tracking.caps"))))
        if self.meta.get("issue"):
            p = self.doc.add_paragraph()
            number = self.sz("issue_number")
            self.fmt(p, number, 0.86, before=30, exact=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            self.text(p, self.t.label("issue") + " ", role="display", size=self.sz("issue_no"), color="muted",
                      bold=False, italic=True)
            self.text(p, str(self.meta["issue"]), role="display", size=number, color="accent", bold=False)
        p = self.doc.add_paragraph(style="Title")
        p.paragraph_format.space_before = Pt(30 if self.meta.get("issue") else self.space("cover_title_top"))
        self._title_lines(p, "title", self.sz("title"))
        lede = self.meta.get("lede") or self.meta.get("subtitle")
        if lede:
            p = self.doc.add_paragraph(style="Lede")
            self.fmt(p, self.sz("subtitle"), self.lead("lede"), before=18, align=WD_ALIGN_PARAGRAPH.LEFT,
                     right=max(0, self.text_width - 330))
            self.text(p, str(lede), role="body", size=self.sz("subtitle"), color="muted")
        self._short_bar(180, 1.5, "accent", before=24)
        self._cover_footer_toc(section, chapters)

    def _cover_footer_toc(self, section, chapters: List[Dict[str, str]]) -> None:
        footer = section.footer
        footer.is_linked_to_previous = False
        if not chapters:
            return
        label = footer.paragraphs[0]
        self.fmt(label, self.sz("toc_label"), 1.3, after=7, align=WD_ALIGN_PARAGRAPH.LEFT)
        self.text(label, self.t.label("toc"), role="label", size=self.sz("toc_label"), color="muted",
                  track=0.1)
        widths = [33, self.text_width - 63, 30]
        table = footer.add_table(rows=len(chapters), cols=3, width=Pt(self.text_width))
        self._plain_table(table, widths)
        for i, chapter in enumerate(chapters):
            cells = [table.cell(i, j) for j in range(3)]
            for cell in cells:
                ox.cell_margins(cell, 8, 0, 8, 0)
                sides = {"top": (0.75, self.t.hex("rule"))}
                if i == len(chapters) - 1:
                    sides["bottom"] = (0.75, self.t.hex("rule"))
                ox.cell_borders(cell, **sides)
            p = cells[0].paragraphs[0]
            self.fmt(p, self.sz("toc_number"), 1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            self.text(p, chapter["number"], role="display", size=self.sz("toc_number"), color="accent", bold=False)
            p = cells[1].paragraphs[0]
            self.fmt(p, self.sz("toc_title"), 1.4, align=WD_ALIGN_PARAGRAPH.LEFT)
            self.text(p, chapter["title"], role="heading", size=self.sz("toc_title"))
            if chapter["desc"]:
                self.text(p, "　" + chapter["desc"], role="body", size=self.sz("toc_desc"), color="muted",
                          bold=False)
            p = cells[2].paragraphs[0]
            self.fmt(p, self.sz("toc_title"), 1.4, align=WD_ALIGN_PARAGRAPH.RIGHT)
            ox.add_field(p, f"PAGEREF {chapter['bookmark']} \\h", "", dirty=True,
                         style=lambda r: self.run(r, "display", self.sz("toc_title"), bold=False))
        footer._element.append(OxmlElement("w:p"))  # 页脚须以段落结尾

    # ------------------------------------------------------------ 正文

    def render(self, blocks: List[Any]):
        cover = self.layout["cover"]
        if self.meta.get("cover") is False:
            cover = "title-block"
        chapters = self._chapters(blocks) if self.layout["chapter"] == "opener" else []
        if cover == "title-block":
            self._chrome(self.doc.sections[0], plain_first_page=True)
            self._title_block()
        else:
            cover_section = self.doc.sections[0]
            if cover == "issue":
                self._cover_issue(cover_section, chapters)
            else:
                self._cover_standard(cover_section)
            body = self.doc.add_section(WD_SECTION.NEW_PAGE)
            self._quiet_break()
            self._chrome(body)
            self._fresh_page = True
        items = attach_captions(blocks)
        i = 0
        while i < len(items):
            block, caption = items[i]
            following = items[i + 1][0] if i + 1 < len(items) else None
            if isinstance(block, mdb.Heading) and block.level == 1 and self.layout["chapter"] == "opener":
                lede = following if isinstance(following, mdb.Quote) and not following.alert else None
                self._chapter_opener(block, chapters, lede)
                i += 2 if lede else 1
                continue
            self._block(block, 0, caption)
            i += 1
        return self.doc

    def _chapters(self, blocks: List[Any]) -> List[Dict[str, str]]:
        chapters = []
        for block in blocks:
            if isinstance(block, mdb.Heading) and block.level == 1:
                text = mdb.plain(block.spans).replace("\n", " ")
                label = block.attrs.get("label", "")
                title, desc = (label.split(" · ")[0], text) if label else (text, "")
                n = len(chapters) + 1
                chapters.append({"number": f"{n:02d}", "title": title, "desc": desc,
                                 "bookmark": f"_acks_chapter_{n}"})
        return chapters

    def _quiet_break(self) -> None:
        """分节符所在的空段落尽量不占高度。"""
        body = self.doc.element.body
        last = body.findall(qn("w:p"))[-1]
        ppr = last.get_or_add_pPr()
        ox.set_child(ppr, ox.el("w:spacing", before=0, after=0, line=20, lineRule="exact"), "pPr")
        rpr = ppr.find(qn("w:rPr"))
        if rpr is None:
            rpr = ox.set_child(ppr, ox.el("w:rPr"), "pPr")
        ox.set_child(rpr, ox.el("w:sz", val=2), "rPr")

    def _set_columns(self, count: int) -> None:
        if count == self.current_cols:
            return
        section = self.doc.add_section(WD_SECTION.CONTINUOUS)
        self._quiet_break()
        ox.columns(section, count, float(self.layout["column_gap_cm"]) * 28.3465)
        self.current_cols = count

    def _chapter_opener(self, heading: mdb.Heading, chapters: List[Dict[str, str]],
                        lede: Optional[mdb.Quote]) -> None:
        """folio 的章节页：另起一页，大号编号、标签、标题、导语、细线，之后正文分栏。"""
        self.chapter_no += 1
        info = chapters[self.chapter_no - 1]
        label = heading.attrs.get("label", "")
        if self._fresh_page:
            section = self.doc.sections[-1]
            if self.current_cols != 1:
                self._set_columns(1)
        else:
            section = self.doc.add_section(WD_SECTION.NEW_PAGE)
            self._quiet_break()
            ox.columns(section, 1, 0)
            self.current_cols = 1
        self._chrome(section, chapter=label.split(" · ")[0] if label else info["desc"] or info["title"])
        self._fresh_page = False
        p = self.doc.add_paragraph()
        number = self.sz("chapter_number")
        self.fmt(p, number, 0.9, exact=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.keep_with_next = True
        self.text(p, info["number"], role="display", size=number, color="accent", bold=False)
        if label:
            head, _, tail = label.partition(" · ")
            self.text(p, "　" + head, role="label", size=self.sz("chapter_label"), track=0.1)
            if tail:
                self.text(p, " · " + tail, role="display", size=self.sz("chapter_label") * 0.92, color="muted",
                          bold=True, track=self.t.get("tracking.caps"))
        title = self.doc.add_paragraph(style="Heading 1")
        title.paragraph_format.space_before = Pt(14)
        _add_spans(title, heading.spans, self._span_style("heading", self.sz("h1"), heading=True))
        ox.bookmark(title, info["bookmark"], next(self._ids))
        if lede is not None:
            text = "\n".join(mdb.plain(b.spans) for b in lede.blocks if isinstance(b, (mdb.Paragraph, mdb.Heading)))
            p = self.doc.add_paragraph(text, style="Lede")
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.right_indent = Pt(max(0, self.text_width - 390))
        rule = self.doc.add_paragraph()
        self.fmt(rule, 2, 1.0, before=16, after=16, exact=True)
        ox.paragraph_borders(rule, bottom=(0.75, self.t.hex("rule"), 0))
        self._set_columns(self.columns)
        self._drop_cap_next = bool(self.layout["drop_cap"])

    def _span_style(self, role: str, size: float, color: str = "ink", heading: bool = False):
        def style(run, span: mdb.Span):
            if span.code:
                self.run(run, "code", size * 0.92, "accent" if not heading else color, bold=False)
            else:
                self.run(run, role, size, color, bold=None if heading else False)
            if span.link:
                run.font.color.rgb = self.rgb("accent")
                run.underline = True
        return style

    def _block(self, block: Any, level: int, caption: Optional[str] = None) -> None:
        full_width = isinstance(block, (mdb.Table, mdb.Code, mdb.Image))
        if self.columns > 1 and self.chapter_no:
            self._set_columns(1 if full_width else self.columns)
        if isinstance(block, mdb.Heading):
            self._heading(block)
        elif isinstance(block, mdb.Paragraph):
            self._paragraph(block.spans)
        elif isinstance(block, mdb.ListBlock):
            self._list(block, level)
        elif isinstance(block, mdb.Table):
            if block.header or block.rows:
                self._table(block, caption)
        elif isinstance(block, mdb.Quote):
            self._quote(block)
        elif isinstance(block, mdb.Code):
            self._code(block.text)
        elif isinstance(block, mdb.Rule):
            p = self.doc.add_paragraph()
            self.fmt(p, 2, 1.0, before=6, after=6, exact=True)
            ox.paragraph_borders(p, bottom=(0.75, self.t.hex("rule"), 0))
        elif isinstance(block, mdb.Image):
            self._image(block)
        self._fresh_page = False

    def _heading(self, block: mdb.Heading) -> None:
        level = min(block.level, 3)
        p = self.doc.add_paragraph(style=f"Heading {level}")
        _add_spans(p, block.spans, self._span_style("heading", self.sz(f"h{level}"), heading=True))

    def _paragraph(self, spans: List[mdb.Span]) -> None:
        if self._drop_cap_next and spans and spans[0].text and not spans[0].code:
            self._drop_cap_next = False
            first, rest = spans[0].text[0], spans[0].text[1:]
            line = self.sz("body") * self.lead("body")
            cap = self.doc.add_paragraph()
            ox.drop_cap(cap, 3, line)
            self.text(cap, first, role="title", size=self.sz("drop_cap"), color="accent")
            spans = ([mdb.Span(rest, spans[0].bold, spans[0].italic, False, spans[0].strike, spans[0].link)]
                     if rest else []) + spans[1:]
        p = self.doc.add_paragraph()
        _add_spans(p, spans, self._span_style("body", self.sz("body")))

    def _list(self, block: mdb.ListBlock, level: int) -> None:
        num_id = self._list_num(block.ordered, min(level, 2), block.start)
        for item in block.items:
            p = self.doc.add_paragraph(style="List Item")
            if item.checked is not None:
                indent = 18 * (level + 1)
                p.paragraph_format.left_indent = Pt(indent)
                p.paragraph_format.first_line_indent = Pt(-14)
                self.text(p, ("☑" if item.checked else "☐") + " ", role="body", size=self.sz("body"),
                          color="accent", bold=False)
            else:
                num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = min(level, 2)
                num_pr.get_or_add_numId().val = num_id
            _add_spans(p, item.spans, self._span_style("body", self.sz("body")))
            for child in item.children:
                self._block(child, level + 1)

    def _plain_table(self, table, widths: List[float]) -> None:
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        ox.table_layout_fixed(table)
        ox.table_no_borders(table)
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Pt(widths[i])
        for i, column in enumerate(table.columns):
            column.width = Pt(widths[i])

    def _caption(self, kind: str, text: str) -> None:
        if kind == "table":
            self.table_no += 1
            number = f"{self.t.label('table')} {self.table_no}"
        else:
            self.figure_no += 1
            number = f"{self.t.label('figure')} {self.figure_no}"
        p = self.doc.add_paragraph(style="Caption")
        size = self.sz("caption")
        if self.layout.get("caption") == "accent":
            self.text(p, number, role="display", size=size * 1.13, color="accent", bold=False, italic=True)
        else:
            self.text(p, number, role="label", size=size, color="muted")
        self.text(p, "　" + text, role="label", size=size, color="muted", bold=False)

    def _table(self, block: mdb.Table, caption: Optional[str]) -> None:
        header = [mdb.plain(c).replace("\n", " ") for c in block.header]
        rows_text = [[mdb.plain(c).replace("\n", " ") for c in r] for r in block.rows]
        width = max([len(header)] + [len(r) for r in rows_text])
        numeric = set(numeric_columns(header, rows_text, block.aligns))
        if caption:
            self._caption("table", caption)
        weights = []
        for i in range(width):
            texts = ([header[i]] if i < len(header) else []) + [r[i] for r in rows_text if i < len(r)]
            weights.append(min(max([display_width(t) for t in texts] + [4]), 40))
        total = sum(weights)
        widths = [self.column_width * w / total for w in weights]
        has_header = bool(header)
        table = self.doc.add_table(rows=len(rows_text) + (1 if has_header else 0), cols=width)
        self._plain_table(table, widths)
        tokens = self.layout["table"]
        cell_v, cell_h = self.space("table_cell_v"), self.space("table_cell_h")
        head_variant = tokens["head"]
        all_rows = ([("head", block.header)] if has_header else []) + [("row", r) for r in block.rows]
        last = len(all_rows) - 1
        for r, (kind, cells) in enumerate(all_rows):
            texts = header if kind == "head" else rows_text[r - (1 if has_header else 0)]
            total_row = kind == "row" and is_total_row(texts)
            ox.row_flags(table.rows[r], header=kind == "head")
            for c in range(width):
                cell = table.cell(r, c)
                ox.cell_margins(cell, cell_v, cell_h, cell_v, cell_h)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.style = self.doc.styles["Table Head" if kind == "head" else "Table Text"]
                if c in numeric:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                spans = cells[c] if c < len(cells) else []
                if kind == "head":
                    size = self.sz("table_head")
                    track = self.t.get("tracking.label") if head_variant == "label" else 0
                    for span in spans:
                        self.text(p, span.text, role="label", size=size, color="muted", bold=False, track=track)
                else:
                    _add_spans(p, spans, self._span_style("body", self.sz("table")))
                    if total_row:
                        for run in p.runs:
                            run.font.bold = True
                sides: Dict[str, Any] = {}
                if kind == "head":
                    if head_variant == "fill":
                        ox.shade_cell(cell, self.t.hex("surface"))
                        sides["top"] = (0.75, self.t.hex(tokens["head_rule"]))
                    else:
                        sides["bottom"] = (1.5, self.t.hex(tokens["head_rule"]))
                elif total_row:
                    # rules：上下各一条细线；bar：上方一条粗线
                    bar = tokens["total"] == "bar"
                    sides["top"] = (1.5 if bar else 0.75, self.t.hex(tokens["total_rule"]))
                    if not bar:
                        sides["bottom"] = (0.75, self.t.hex(tokens["total_rule"]))
                elif r == last:
                    sides["bottom"] = (0.75, self.t.hex(tokens["row_rule" if tokens["total"] == "bar" else "total_rule"]))
                elif not (r + 1 <= last and is_total_row(rows_text[r + 1 - (1 if has_header else 0)])):
                    sides["bottom"] = (0.75, self.t.hex(tokens["row_rule"]))
                if sides:
                    ox.cell_borders(cell, **sides)
        spacer = self.doc.add_paragraph()
        self.fmt(spacer, 4, 1.0, exact=True, after=self.space("block"))

    def _box(self, fill: Optional[str], rules: bool) -> Any:
        table = self.doc.add_table(rows=1, cols=1)
        self._plain_table(table, [self.column_width])
        cell = table.cell(0, 0)
        if fill:
            ox.shade_cell(cell, fill)
        if rules:
            ox.cell_borders(cell, top=(0.75, self.t.hex("rule_strong")), bottom=(0.75, self.t.hex("rule_strong")))
        ox.row_flags(table.rows[0])
        return cell

    def _quote(self, block: mdb.Quote) -> None:
        paragraphs = [b for b in block.blocks if isinstance(b, (mdb.Paragraph, mdb.Heading, mdb.ListBlock))]
        if block.alert:
            self._callout(block.alert, paragraphs)
        elif self.layout["quote"] == "pull":
            self._pull_quote(paragraphs)
        else:
            for b in paragraphs:
                for spans in _spans_of(b):
                    p = self.doc.add_paragraph(style="Quote")
                    p.paragraph_format.left_indent = Pt(14)
                    ox.paragraph_borders(p, left=(2.25, self.t.hex("rule_mid"), 10))
                    _add_spans(p, spans, self._span_style("quote", self.sz("quote"), "muted"))

    def _callout(self, alert: str, blocks: List[Any]) -> None:
        tint = self.layout["callout"] == "tint"
        cell = self._box(self.t.hex("surface") if tint else None, rules=not tint)
        pad_v, pad_h = self.space("callout_pad_v"), self.space("callout_pad_h")
        ox.cell_margins(cell, pad_v, pad_h if tint else 0, pad_v, pad_h if tint else 0)
        label = cell.paragraphs[0]
        size = self.sz("callout_label")
        self.fmt(label, size, 1.3, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
        self.text(label, self.t.label(alert), role="label", size=size, color="accent", bold=True, track=0.1)
        for b in blocks:
            for spans in _spans_of(b):
                p = cell.add_paragraph(style="Callout Text")
                _add_spans(p, spans, self._span_style("body", self.sz("callout")))
        spacer = self.doc.add_paragraph()
        self.fmt(spacer, 4, 1.0, exact=True, after=self.space("block"))

    def _pull_quote(self, blocks: List[Any]) -> None:
        table = self.doc.add_table(rows=1, cols=2)
        size = self.sz("pull_quote")
        mark_width = size * 1.7
        self._plain_table(table, [mark_width, self.column_width - mark_width])
        ox.row_flags(table.rows[0])
        for cell in table.rows[0].cells:
            ox.cell_margins(cell, 10, 0, 10, 0)
            ox.cell_borders(cell, top=(0.75, self.t.hex("rule_strong")), bottom=(0.75, self.t.hex("rule_strong")))
        mark = table.cell(0, 0).paragraphs[0]
        self.fmt(mark, size * 1.8, 0.9, exact=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        self.text(mark, "“", role="display", size=size * 1.8, color="accent", bold=False)
        first = True
        for b in blocks:
            for spans in _spans_of(b):
                cell = table.cell(0, 1)
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                self.fmt(p, size, self.lead("pull_quote"), align=WD_ALIGN_PARAGRAPH.LEFT)
                _add_spans(p, spans, self._span_style("quote", size))
        spacer = self.doc.add_paragraph()
        self.fmt(spacer, 4, 1.0, exact=True, after=self.space("block"))

    def _code(self, text: str) -> None:
        cell = self._box(self.t.hex("surface"), rules=False)
        ox.cell_margins(cell, 8, 10, 8, 10)
        for i, line in enumerate(text.split("\n")):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.style = self.doc.styles["Code"]
            self.text(p, line or " ", role="code", size=self.sz("code"), bold=False)
        spacer = self.doc.add_paragraph()
        self.fmt(spacer, 4, 1.0, exact=True, after=self.space("block"))

    def _image(self, block: mdb.Image) -> None:
        path = block.src if os.path.isabs(block.src) else os.path.join(self.base_dir, block.src)
        if not os.path.isfile(path):
            self.warnings.append({"code": "IMAGE_NOT_FOUND", "message": f"图片不存在：{block.src}"})
            self._paragraph([mdb.Span(f"[图片：{block.alt or block.src}]")])
            return
        shape = self.doc.add_picture(path)
        limit = Emu(int(self.column_width * 12700))
        if shape.width > limit:
            shape.height = int(shape.height * limit / shape.width)
            shape.width = limit
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(2)
        if block.alt:
            self._caption("figure", block.alt)


def _spans_of(block: Any) -> List[List[mdb.Span]]:
    if isinstance(block, (mdb.Paragraph, mdb.Heading)):
        return [block.spans]
    if isinstance(block, mdb.ListBlock):
        return [[mdb.Span("· ")] + item.spans for item in block.items]
    return []


def _add_hyperlink(paragraph, url: str) -> Run:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    link.append(r)
    paragraph._p.append(link)
    return Run(r, paragraph)


def _add_spans(paragraph, spans: List[mdb.Span], style_run) -> None:
    for span in spans:
        run = _add_hyperlink(paragraph, span.link) if span.link else paragraph.add_run()
        run.text = span.text
        style_run(run, span)
        if span.bold:
            run.bold = True
        if span.italic:
            run.italic = True
        if span.strike:
            run.font.strike = True


def render_word(title: Optional[str], content: str, output_path: str, theme: Any = None,
                base_dir: Optional[str] = None, font_policy: str = "local",
                footer_label: Optional[str] = None, **meta: Any) -> Dict[str, Any]:
    """按主题生成 Word。meta 可传 subtitle、kicker、brand、author、date、version、issue、lede 等，
    会覆盖正文开头 front matter 里的同名字段。"""
    theme = load_theme(theme)
    info, body = document_meta(title, content, {k: v for k, v in meta.items() if k != "theme"})
    renderer = WordRenderer(theme, info, base_dir or os.getcwd(), font_policy, footer_label)
    renderer.render(mdb.parse(body))
    renderer.doc.save(output_path)
    result: Dict[str, Any] = {"output_path": output_path, "file_size": os.path.getsize(output_path),
                              "theme": theme.name}
    if renderer.warnings:
        result["warnings"] = renderer.warnings
    return result
