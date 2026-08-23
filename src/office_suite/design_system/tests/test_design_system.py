"""Unit tests for ACKS Studio Design System v2.1."""

import json
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent


# ── Token sync tests ────────────────────────────────────────────────

class TestTokenSync:
    """Verify generated constants match tokens.json."""

    def test_tokens_json_valid(self):
        with open(ROOT / "tokens.json", encoding="utf-8") as f:
            tok = json.load(f)
        assert "color" in tok
        assert "font" in tok
        assert "version" in tok

    def test_primary_color_consistent(self):
        with open(ROOT / "tokens.json", encoding="utf-8") as f:
            tok = json.load(f)
        expected = tok["color"]["brand"]["PRIMARY"].lstrip("#").upper()

        # tokens_constants.py
        tc = (ROOT / "tokens_constants.py").read_text()
        assert f'COLOR_PRIMARY = "{expected}"' in tc

        # slides_constants.py
        sc = (ROOT / "slides_constants.py").read_text()
        r, g, b = int(expected[0:2], 16), int(expected[2:4], 16), int(expected[4:6], 16)
        assert f"RGBColor(0x{r:02X}, 0x{g:02X}, 0x{b:02X})" in sc

    def test_font_consistent(self):
        with open(ROOT / "tokens.json", encoding="utf-8") as f:
            tok = json.load(f)
        expected = tok["font"]["family"]["DISPLAY_EN"]

        tc = (ROOT / "tokens_constants.py").read_text()
        assert f'FONT_DISPLAY_EN = "{expected}"' in tc


# ── Module import tests ─────────────────────────────────────────────

class TestImports:
    """Verify modules import and expose __all__."""

    def test_tokens_import(self):
        import tokens
        assert hasattr(tokens, "__all__")
        assert len(tokens.__all__) > 0

    def test_slides_import(self):
        import slides
        assert hasattr(slides, "__all__")
        assert len(slides.__all__) > 0

    def test_xlsx_import(self):
        import xlsx
        assert hasattr(xlsx, "__all__")
        assert len(xlsx.__all__) > 0


# ── FontStyle tests ─────────────────────────────────────────────────

class TestFontStyle:
    """Verify FontStyle dataclass."""

    def test_create_default(self):
        from tokens import FontStyle
        s = FontStyle()
        assert s.font == "DM Sans"
        assert s.bold is False

    def test_frozen(self):
        from tokens import FontStyle
        s = FontStyle()
        with pytest.raises(AttributeError):
            s.font = "Other"

    def test_presets_exist(self):
        from tokens import STYLE_HEADING, STYLE_BODY, STYLE_META
        from docx.shared import Pt
        assert STYLE_HEADING.bold is True
        assert STYLE_BODY.size == Pt(10.5)


# ── Document generation tests ───────────────────────────────────────

class TestDocxGeneration:
    """Verify python-docx generation."""

    def test_build_document(self):
        from tokens import build_document
        doc = build_document()
        assert len(doc.paragraphs) > 0

    def test_save_demo(self, tmp_path):
        from tokens import save_demo
        out = save_demo(str(tmp_path / "test.docx"))
        assert Path(out).exists()
        assert Path(out).stat().st_size > 0

    def test_custom_kpis(self):
        from tokens import build_document
        kpis = [{"label": "Test", "value": "100", "delta": "+10%"}]
        doc = build_document(kpis=kpis)
        assert len(doc.paragraphs) > 0


# ── Presentation generation tests ───────────────────────────────────

class TestPptxGeneration:
    """Verify python-pptx generation."""

    def test_build_presentation(self):
        from slides import build_presentation
        prs = build_presentation()
        assert len(prs.slides) == 11

    def test_save_demo(self, tmp_path):
        from slides import save_demo
        out = save_demo(str(tmp_path / "test.pptx"))
        assert Path(out).exists()
        assert Path(out).stat().st_size > 0

    def test_custom_data(self):
        from slides import build_presentation
        kpis = [{"label": "Test", "value": "100", "delta": "+10%"}]
        prs = build_presentation(kpis=kpis)
        assert len(prs.slides) == 11

    def test_footers_finalized_with_real_total(self):
        """finalize_footers() must patch every deferred '??' placeholder to
        the deck's real slide count, not the old hardcoded default of 8."""
        from slides import build_presentation
        prs = build_presentation()
        total = len(prs.slides)
        found_footer = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.name == "acks_footer_total":
                    found_footer = True
                    text = "".join(r.text for p in shape.text_frame.paragraphs for r in p.runs)
                    assert "??" not in text
                    assert f"{total:02d}" in text
        assert found_footer, "expected at least one acks_footer_total shape"


# ── New v2.1 component tests ────────────────────────────────────────

class TestFitHelpers:
    """Verify the real font-metric auto-fit helpers."""

    def test_fit_title_size_shrinks_long_text(self):
        from slides import fit_title_size
        short = fit_title_size("Hi", 11.6, base_pt=92, min_pt=52)
        long = fit_title_size(
            "A Much Longer Title That Should Not Fit At Full Size", 11.6,
            base_pt=92, min_pt=52,
        )
        assert long <= short
        assert long >= 52

    def test_fit_paragraph_size_respects_min(self):
        from slides import fit_paragraph_size
        size = fit_paragraph_size(
            "word " * 500, box_width_in=4.0, box_height_in=1.0,
            base_pt=18, min_pt=10,
        )
        assert 10 <= size <= 18


class TestPptxNewComponents:
    """Verify the v2.1 list/checklist/template slide components render."""

    def test_add_list_slide(self):
        from slides import init_presentation, add_list_slide
        prs = init_presentation()
        add_list_slide(prs, eyebrow="SECTION", title_en="Title", title_zh="标题",
                        items=["one", "two", "three"], page_no=1)
        assert len(prs.slides) == 1

    def test_add_checklist_slide(self):
        from slides import init_presentation, add_checklist_slide
        prs = init_presentation()
        add_checklist_slide(prs, eyebrow="SECTION", title_en="Title", title_zh="标题",
                             items=["item a", "item b"], page_no=1)
        assert len(prs.slides) == 1

    def test_add_template_slide(self):
        from slides import init_presentation, add_template_slide
        prs = init_presentation()
        add_template_slide(prs, eyebrow="SECTION", title_en="Title", title_zh="标题",
                            code="Some template body text.", page_no=1)
        assert len(prs.slides) == 1


class TestDocxNewComponents:
    """Verify the v2.1 bullet/numbered/checklist/code-block components render."""

    def test_add_bullet_list(self):
        from tokens import build_document, add_bullet_list
        doc = build_document()
        n_before = len(doc.paragraphs)
        add_bullet_list(doc, ["a", "b", "c"])
        assert len(doc.paragraphs) == n_before + 3

    def test_add_numbered_list(self):
        from tokens import build_document, add_numbered_list
        doc = build_document()
        n_before = len(doc.paragraphs)
        add_numbered_list(doc, ["a", "b"])
        assert len(doc.paragraphs) == n_before + 2

    def test_add_checklist(self):
        from tokens import build_document, add_checklist
        doc = build_document()
        n_before = len(doc.paragraphs)
        add_checklist(doc, ["task a", "task b"])
        assert len(doc.paragraphs) == n_before + 2

    def test_add_code_block(self):
        from tokens import build_document, add_code_block
        doc = build_document()
        n_before = len(doc.paragraphs)
        add_code_block(doc, "print('hi')", caption="example")
        assert len(doc.paragraphs) > n_before


# ── Geometry / overflow tests ───────────────────────────────────────

class TestSlideGeometry:
    """Static bounds check: no shape should fall outside the slide or
    bleed into the footer-safe zone. Catches the class of title/insight
    overflow bugs fixed in v2.1 without needing a full LibreOffice render."""

    def test_no_shape_out_of_bounds_or_footer_overflow(self):
        from pptx.util import Emu
        from slides import build_presentation, FOOTER_SAFE_TOP

        prs = build_presentation()
        slide_w, slide_h = prs.slide_width, prs.slide_height
        tolerance = Emu(9144)  # ~0.01in
        footer_pad = Emu(int(0.35 * 914400))

        problems = []
        for i, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                l, t, w, h = shape.left, shape.top, shape.width, shape.height
                if None in (l, t, w, h):
                    continue
                right, bottom = l + w, t + h
                if l < 0 or t < 0:
                    problems.append(f"slide {i}: {shape.name!r} negative origin")
                if right > slide_w + tolerance:
                    problems.append(f"slide {i}: {shape.name!r} exceeds slide width")
                if bottom > slide_h + tolerance:
                    problems.append(f"slide {i}: {shape.name!r} exceeds slide height")
                is_footer = (shape.name or "").startswith("acks_footer")
                if not is_footer and t < FOOTER_SAFE_TOP and bottom > FOOTER_SAFE_TOP + footer_pad:
                    problems.append(f"slide {i}: {shape.name!r} bleeds into footer-safe zone")

        assert not problems, "\n".join(problems)


# ── Spreadsheet generation tests ────────────────────────────────────

class TestXlsxGeneration:
    """Verify openpyxl generation."""

    def test_build_workbook(self):
        from xlsx import build_workbook
        wb = build_workbook()
        sheet_names = wb.sheetnames
        assert "README" in sheet_names
        assert "Dashboard" in sheet_names
        assert "Data" in sheet_names
        assert "Model" in sheet_names

    def test_save_demo(self, tmp_path):
        from xlsx import save_demo
        out = save_demo(str(tmp_path / "test.xlsx"))
        assert Path(out).exists()
        assert Path(out).stat().st_size > 0

    def test_named_styles_registered(self):
        from xlsx import build_workbook
        wb = build_workbook()
        style_names = {s.name if hasattr(s, 'name') else s for s in wb.named_styles}
        assert "ACKS_CELL_H1" in style_names
        assert "ACKS_CELL_KPI_VALUE" in style_names
