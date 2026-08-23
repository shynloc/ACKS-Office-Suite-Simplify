import os
import re
from typing import Optional, List, Dict, Any
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_PARAGRAPH_ALIGNMENT
from docx.enum.section import WD_SECTION


def _is_cjk(text: str) -> bool:
    """Return True if text contains CJK characters."""
    return any('\u4e00' <= ch <= '\u9fff' for ch in text)


def create_word(title: str, content: str, output_path: str,
                theme: str = "acks",
                brand_name: str = "ACKS Studio",
                footer_label: Optional[str] = None,
                **kwargs) -> Dict[str, Any]:
    """
    创建 Word 文档。

    Args:
        title: 文档标题
        content: 文档内容，支持 Markdown 格式的标题、列表
        output_path: 输出文件路径
        theme: "acks"（符合 ACKS 设计规范，默认）或 "default"（简单样式）
        brand_name: 封面品牌 stamp（theme="acks" 生效），传 "" 去品牌化
        footer_label: 页脚 label，默认自动生成
    """
    if theme == "acks":
        return _create_word_acks(title, content, output_path,
                                 brand_name, footer_label)
    return _create_word_default(title, content, output_path)


def _create_word_acks(title: str, content: str, output_path: str,
                      brand_name: str, footer_label: Optional[str]) -> Dict[str, Any]:
    """用 ACKS 设计规范（design_system.tokens）生成 Word 文档。"""
    from .design_system import tokens as acks

    doc = Document()
    acks.apply_default_styles(doc)
    acks.set_a4_page_margins(doc)

    label = footer_label or (f"{brand_name} · 文档设计规范 v2" if brand_name else "文档设计规范 v2")
    acks.add_page_footer(doc, label=label)

    _add_cover_acks(doc, title, brand_name, acks)
    _render_content_acks(doc, content, acks)

    doc.save(output_path)
    return {
        "output_path": output_path,
        "file_size": os.path.getsize(output_path),
        "pages": int(len(doc.paragraphs) / 30) + 1,  # 估算页数
        "theme": "acks",
    }


def _add_cover_acks(doc, title: str, brand_name: str, acks) -> None:
    """根据标题语言映射到 ACKS 封面组件。"""
    if _is_cjk(title):
        acks.add_cover_page(
            doc,
            title_top="", title_em=title, zh_title=title,
            doctype="", brand_name=brand_name,
        )
    else:
        acks.add_cover_page(
            doc,
            title_top=title, title_em="", zh_title="",
            doctype="", brand_name=brand_name,
        )


def _render_content_acks(doc, content: str, acks) -> None:
    """把 Markdown 内容渲染为 ACKS 组件（章节标题/正文/列表）。"""
    lines = content.split('\n')
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # 项目符号列表（聚合连续项）
        if line.startswith('- ') or line.startswith('* '):
            items = []
            while i < n and (lines[i].strip().startswith('- ') or lines[i].strip().startswith('* ')):
                items.append(lines[i].strip()[2:])
                i += 1
            acks.add_bullet_list(doc, items)
            continue
        # 编号列表（聚合连续项）
        m = re.match(r'^(\d+)\.\s+(.*)$', line.strip())
        if m:
            items = []
            while i < n:
                m2 = re.match(r'^(\d+)\.\s+(.*)$', lines[i].strip())
                if not m2:
                    break
                items.append(m2.group(2))
                i += 1
            acks.add_numbered_list(doc, items)
            continue
        # 标题 / 正文
        if line.startswith('# '):
            acks.add_section_heading(doc, en=line[2:])
        elif line.startswith('## '):
            acks.add_sub_heading(doc, en=line[3:])
        elif line.startswith('### '):
            acks.add_label(doc, line[4:])
        else:
            acks.add_body_paragraph(doc, line)
        i += 1


def _create_word_default(title: str, content: str, output_path: str) -> Dict[str, Any]:
    """原简单样式（宋体 + 蓝色标题）。"""
    doc = Document()

    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal'].font.size = Pt(12)

    title_para = doc.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.runs[0].font.size = Pt(24)
    title_para.runs[0].font.bold = True
    title_para.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph()

    lines = content.split('\n')
    for line in lines:
        line = line.rstrip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith('# '):
            para = doc.add_heading(line[2:], level=1)
            para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = RGBColor(0, 102, 204)
        elif line.startswith('## '):
            para = doc.add_heading(line[3:], level=2)
            para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = RGBColor(51, 153, 255)
        elif line.startswith('### '):
            para = doc.add_heading(line[4:], level=3)
            para.runs[0].font.bold = True
        elif line.startswith('- ') or line.startswith('* '):
            para = doc.add_paragraph(style='List Bullet')
            para.add_run(line[2:])
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
            para = doc.add_paragraph(style='List Number')
            para.add_run(line)
        else:
            para = doc.add_paragraph(line)

    doc.save(output_path)

    return {
        "output_path": output_path,
        "file_size": os.path.getsize(output_path),
        "pages": int(len(doc.paragraphs) / 30) + 1,
        "theme": "default",
    }


def add_watermark(input_path: str, watermark_text: str, output_path: Optional[str] = None,
                  font_size: int = 48, color: tuple = (200, 200, 200),
                  italic: bool = True, **kwargs) -> Dict[str, Any]:
    """
    给 Word 文档添加水印（页眉斜体灰色大字，简化版）。

    Args:
        input_path: 输入文件路径
        watermark_text: 水印文本
        output_path: 输出文件路径（默认覆盖输入）
        font_size: 字号
        color: RGB 颜色元组
        italic: 是否斜体
    """
    if not output_path:
        output_path = input_path

    doc = Document(input_path)

    section = doc.sections[0]
    header = section.header

    watermark_para = header.add_paragraph()
    watermark_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = watermark_para.add_run(watermark_text)
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color)
    run.bold = True
    run.italic = italic

    doc.save(output_path)

    return {"output_path": output_path}


def extract_text(input_path: str, **kwargs) -> str:
    """
    提取 Word 文档中的文本
    """
    doc = Document(input_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)


def merge_documents(input_paths: List[str], output_path: str, **kwargs) -> Dict[str, Any]:
    """
    合并多个 Word 文档
    """
    merged_doc = Document()

    for i, path in enumerate(input_paths):
        if not os.path.exists(path):
            continue
        if i > 0:
            merged_doc.add_section(WD_SECTION.NEW_PAGE)
        source_doc = Document(path)
        for element in source_doc.element.body:
            merged_doc.element.body.append(element)

    merged_doc.save(output_path)
    return {"output_path": output_path, "merged_count": len(input_paths)}
